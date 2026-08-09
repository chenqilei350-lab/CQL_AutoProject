"""Split the approved combined report into code-first and experiment-only Word files."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "reports" / "AUT_KG_Complete_Experiment_and_Code_Report_ZH_2026-08-03.docx"
CODE_OUTPUT = ROOT / "docs" / "reports" / "AUT_KG_Code_Module_Detailed_Report_ZH_2026-08-04.docx"
EXPERIMENT_OUTPUT = ROOT / "docs" / "reports" / "AUT_KG_Experiment_Module_Report_ZH_2026-08-04.docx"


def _body_paragraph_text(document: Document) -> dict[int, str]:
    body = document.element.body
    texts: dict[int, str] = {}
    for index, element in enumerate(list(body)):
        if element.tag == qn("w:p"):
            texts[index] = Paragraph(element, document._body).text.strip()
    return texts


def _find_heading(texts: dict[int, str], expected: str) -> int:
    for index, text in texts.items():
        if text == expected:
            return index
    raise ValueError(f"Heading not found in source document: {expected}")


def _trim_document(document: Document, keep_indices: set[int]) -> None:
    body = document.element.body
    for index, element in reversed(list(enumerate(list(body)))):
        if element.tag == qn("w:sectPr"):
            continue
        if index not in keep_indices:
            body.remove(element)


def _replace_paragraph_text(document: Document, replacements: dict[str, str]) -> None:
    for paragraph in document.paragraphs:
        stripped = paragraph.text.strip()
        replacement = replacements.get(stripped)
        if replacement is not None:
            paragraph.text = replacement


def _replace_table_text(document: Document, old: str, new: str) -> None:
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() == old:
                    cell.text = new


def _set_header(document: Document, text: str) -> None:
    for section in document.sections:
        if section.header.paragraphs:
            section.header.paragraphs[0].text = text


def _base_indices(document: Document) -> tuple[dict[int, str], dict[str, int], int]:
    texts = _body_paragraph_text(document)
    markers = {
        "summary": _find_heading(texts, "执行摘要"),
        "code": _find_heading(texts, "10. 代码架构与模块实现"),
        "issues": _find_heading(texts, "11. 遇到的问题与对应改进"),
        "appendix_d": _find_heading(texts, "附录 D：代码、方法与数据来源索引"),
    }
    body_length = len(list(document.element.body))
    return texts, markers, body_length


def build_code_report() -> Path:
    document = Document(SOURCE)
    _, markers, body_length = _base_indices(document)
    keep = set(range(0, markers["summary"]))
    keep.update(range(markers["code"], markers["issues"]))
    keep.update(range(markers["appendix_d"], body_length))
    _trim_document(document, keep)

    replacements = {
        "AUT KG Extraction Pipeline\n实验与代码全量总结报告": "AUT KG Extraction Pipeline\n代码模块详细报告",
        "代码架构、关键实现、实验结果、问题与改进措施": "代码架构、模块职责、关键实现、实验对应效果与来源",
        "10. 代码架构与模块实现": "1. 代码架构与模块实现",
        "10.1 总体代码架构": "1.1 总体代码架构",
        "10.2 Exposé 原计划与当前代码方法对照": "1.2 Exposé 原计划与当前代码方法对照",
        "10.2.1 Exposé 中原计划采用的方法": "1.2.1 Exposé 中原计划采用的方法",
        "10.2.2 当前代码实际采用的方法": "1.2.2 当前代码实际采用的方法",
        "10.2.3 原计划与当前实现的逐项差异": "1.2.3 原计划与当前实现的逐项差异",
        "10.2.4 为什么采用当前的阶段性实现": "1.2.4 为什么采用当前的阶段性实现",
        "10.3 数据清洗模块：确定性来源解析": "1.3 数据清洗模块：确定性来源解析",
        "10.4 统一文本模块：只重组原文支持事实": "1.4 统一文本模块：只重组原文支持事实",
        "10.5 Schema 与本地结构化 LLM 抽取": "1.5 Schema 与本地结构化 LLM 抽取",
        "10.6 实验编排与逐层抽取": "1.6 实验编排与逐层抽取",
        "10.7 关系候选生成与二元判断": "1.7 关系候选生成与二元判断",
        "10.8 Property Graph 构建": "1.8 Property Graph 构建",
        "10.9 工业后处理与关系来源追踪": "1.9 工业后处理与关系来源追踪",
        "10.10 评价：P/R/F1、Hallucination 与 Stability": "1.10 评价：P/R/F1、Hallucination 与 Stability",
        "10.11 Gold 与公开数据 Adapter": "1.11 Gold 与公开数据 Adapter",
        "10.12 一键产品入口与输出格式": "1.12 一键产品入口与输出格式",
        "10.13 代码模块与实验结论对照": "1.13 代码模块与实验结论对照",
        "附录 D：代码、方法与数据来源索引": "附录：代码、方法与数据来源索引",
        "D.1 直接代码依赖与 API": "A.1 直接代码依赖与 API",
        "D.2 方法与系统设计参考": "A.2 方法与系统设计参考",
        "D.3 数据集与标注来源": "A.3 数据集与标注来源",
    }
    _replace_paragraph_text(document, replacements)
    for paragraph in document.paragraphs:
        if paragraph.text.startswith("以下来源编号与第 10 章一致"):
            paragraph.text = paragraph.text.replace("第 10 章", "第 1 章", 1)
    _replace_table_text(document, "2026-08-03", "2026-08-04")
    _set_header(document, "AUT KG Extraction Pipeline  |  Code Module Report")
    document.core_properties.title = "AUT KG Extraction Pipeline - Code Module Detailed Report"
    CODE_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(CODE_OUTPUT)
    return CODE_OUTPUT


def build_experiment_report() -> Path:
    document = Document(SOURCE)
    _, markers, _ = _base_indices(document)
    keep = set(range(0, markers["code"]))
    keep.update(range(markers["issues"], markers["appendix_d"]))
    _trim_document(document, keep)

    replacements = {
        "AUT KG Extraction Pipeline\n实验与代码全量总结报告": "AUT KG Extraction Pipeline\n实验模块全量报告",
        "代码架构、关键实现、实验结果、问题与改进措施": "实验目的、实现方式、评价标准、结果、问题与改进措施",
        "11. 遇到的问题与对应改进": "10. 遇到的问题与对应改进",
        "11.1 改进措施的方法来源": "10.1 改进措施的方法来源",
        "12. 模块保留决策": "11. 模块保留决策",
        "12.1 当前真正完成的系统": "11.1 当前真正完成的系统",
        "13. 最终研究成果与下一步": "12. 最终研究成果与下一步",
        "13.1 已经得到的可靠结论": "12.1 已经得到的可靠结论",
        "13.2 优先级计划": "12.2 优先级计划",
        "13.3 可用于论文的最终表述": "12.3 可用于论文的最终表述",
    }
    _replace_paragraph_text(document, replacements)
    _replace_table_text(document, "2026-08-03", "2026-08-04")
    _set_header(document, "AUT KG Extraction Pipeline  |  Experiment Module Report")
    document.core_properties.title = "AUT KG Extraction Pipeline - Experiment Module Report"
    EXPERIMENT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(EXPERIMENT_OUTPUT)
    return EXPERIMENT_OUTPUT


if __name__ == "__main__":
    print(build_code_report())
    print(build_experiment_report())
