#!/usr/bin/env python3
"""Build the complete Chinese experiment report as a verified Word document."""

from __future__ import annotations

import io
import keyword
from pathlib import Path
import token
import tokenize
import unicodedata

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
BUILD_DIR = ROOT / "tmp" / "experiment_report_20260803"
OUTPUT = ROOT / "docs" / "reports" / "AUT_KG_Complete_Experiment_and_Code_Report_ZH_2026-08-03.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17212B"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_RED = "FDECEC"
RED = "9B1C1C"
GREEN = "1F5F3B"
WHITE = "FFFFFF"
CONTENT_WIDTH_DXA = 9360
CJK_FONT = "Noto Sans SC"
CODE_ASCII_FONT = Path("/System/Library/Fonts/Menlo.ttc")
CODE_CJK_FONT = BUILD_DIR / "fonts" / "NotoSansSC-Regular.ttf"
CODE_CJK_BOLD_FONT = BUILD_DIR / "fonts" / "NotoSansSC-Bold.ttf"


def set_run_font(
    run,
    *,
    name: str = CJK_FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def style_table_borders(table, color: str = "B7C3D0") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_paragraph_border_bottom(paragraph, color: str = BLUE, size: int = 10) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_paragraph(paragraph, fill: str, left_border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if left_border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        borders.append(left)
        p_pr.append(borders)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    end = paragraph.add_run(" 页")
    set_run_font(end, size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = CJK_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), CJK_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), CJK_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    style_values = {
        "Title": (24, INK, 0, 6),
        "Subtitle": (13, MUTED, 0, 14),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in style_values.items():
        style = doc.styles[style_name]
        style.font.name = CJK_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), CJK_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), CJK_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for current in doc.sections:
        header_p = current.header.paragraphs[0]
        header_p.text = "AUT KG Extraction Pipeline  |  Experiment Report"
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(header_p.runs[0], size=8.5, color=MUTED, bold=True)
        set_paragraph_border_bottom(header_p, color="D5DDE6", size=4)
        add_page_number(current.footer.paragraphs[0])


def add_body(doc: Document, text: str, *, bold_lead: str | None = None, after: float = 6) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)


def add_labeled(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    lead = p.add_run(f"{label}：")
    set_run_font(lead, bold=True, color=DARK_BLUE)
    run = p.add_run(text)
    set_run_font(run)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.clear()
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.clear()
        run = p.add_run(item)
        set_run_font(run)


def add_callout(doc: Document, title: str, text: str, *, risk: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    shade_paragraph(p, PALE_RED if risk else LIGHT_BLUE, RED if risk else BLUE)
    lead = p.add_run(f"{title}  ")
    set_run_font(lead, bold=True, color=RED if risk else DARK_BLUE)
    body = p.add_run(text)
    set_run_font(body)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[object]],
    widths: list[int],
    *,
    font_size: float = 8.5,
    first_column_bold: bool = False,
    zebra: bool = True,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    style_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=font_size, bold=True, color=DARK_BLUE)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.text = ""
            if zebra and row_index % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if col_index > 0 and str(value).replace(".", "", 1).replace("-", "", 1).isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run_font(
                run,
                size=font_size,
                bold=first_column_bold and col_index == 0,
                color=INK,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"数据来源：{text}")
    set_run_font(run, size=8.5, color=MUTED, italic=True)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """Append a compact external hyperlink to a paragraph."""

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), CJK_FONT)
    fonts.set(qn("w:hAnsi"), CJK_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_properties, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_reference_entry(doc: Document, source_id: str, title: str, url: str, use: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    lead = paragraph.add_run(f"[{source_id}] {title}。")
    set_run_font(lead, bold=True, color=DARK_BLUE)
    body = paragraph.add_run(f"用于：{use}。来源：")
    set_run_font(body)
    add_hyperlink(paragraph, "打开官方页面", url)


def add_figure(doc: Document, path: Path, caption: str, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    picture = run.add_picture(str(path), width=Inches(6.25))
    picture._inline.docPr.set("descr", alt)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(8)
    cap = caption_p.add_run(caption)
    set_run_font(cap, size=9, color=MUTED, italic=True)


def add_code_figure(doc: Document, path: Path, caption: str, alt: str) -> None:
    """Add a code screenshot with enough width for readable line numbers."""

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    picture = run.add_picture(str(path), width=Inches(6.2))
    picture._inline.docPr.set("descr", alt)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(7)
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=8.5, color=MUTED, italic=True)


def _code_font(size: int, *, cjk: bool = False, bold: bool = False):
    if cjk:
        candidates = [
            CODE_CJK_BOLD_FONT if bold else CODE_CJK_FONT,
            Path(
                "/Applications/Goodnotes.app/Contents/Resources/"
                "GN6App_GN6AppServices.bundle/Contents/Resources/"
                f"NotoSansSC-{'Bold' if bold else 'Regular'}.ttf"
            ),
            Path("/System/Library/Fonts/PingFang.ttc"),
        ]
    else:
        candidates = [CODE_ASCII_FONT, Path("/System/Library/Fonts/Monaco.ttf")]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _code_token_style(token_type: int, token_text: str) -> tuple[str, bool]:
    if token_type == token.COMMENT:
        return "#7A8490", False
    if token_type == token.STRING:
        return "#477A4B", False
    if token_type == token.NUMBER:
        return "#B05A00", False
    if token_type == token.OP:
        return "#5A6570", False
    if token_type == token.NAME and keyword.iskeyword(token_text):
        return "#6F42C1", True
    if token_type == token.NAME and token_text in {
        "BaseModel",
        "Field",
        "Path",
        "Protocol",
        "PropertyGraph",
        "EgocentricVideoExtraction",
    }:
        return "#1769AA", False
    return "#17212B", False


def _display_cells(text: str) -> int:
    cells = 0
    for character in text.expandtabs(4):
        cells += 2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1
    return cells


def draw_code_excerpt(
    output_path: Path,
    source_path: Path,
    start_line: int,
    end_line: int,
) -> None:
    """Render a repository excerpt as a line-numbered PNG without altering code."""

    all_lines = source_path.read_text(encoding="utf-8").splitlines()
    selected = all_lines[start_line - 1 : end_line]
    code = "\n".join(selected) + "\n"
    styles: list[list[tuple[str, bool]]] = [
        [("#17212B", False) for _ in line]
        for line in selected
    ]
    try:
        tokens = tokenize.generate_tokens(io.StringIO(code).readline)
        for token_info in tokens:
            color, bold = _code_token_style(token_info.type, token_info.string)
            start_row, start_column = token_info.start
            end_row, end_column = token_info.end
            for row in range(start_row, end_row + 1):
                if row < 1 or row > len(selected):
                    continue
                first_column = start_column if row == start_row else 0
                last_column = end_column if row == end_row else len(selected[row - 1])
                for column in range(first_column, min(last_column, len(selected[row - 1]))):
                    styles[row - 1][column] = (color, bold)
    except (IndentationError, tokenize.TokenError):
        pass

    font_size = 21
    cell_width = 13
    line_height = 31
    header_height = 56
    number_width = 90
    horizontal_padding = 24
    max_cells = max((_display_cells(line) for line in selected), default=1)
    width = max(1200, min(2600, number_width + horizontal_padding * 2 + max_cells * cell_width))
    height = header_height + 22 + max(1, len(selected)) * line_height + 24
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, header_height), fill="#263746")
    draw.rectangle((0, header_height, number_width, height), fill="#F1F4F7")
    draw.line((number_width, header_height, number_width, height), fill="#D6DEE6", width=2)

    header_font = _code_font(20, cjk=True, bold=True)
    line_number_font = _code_font(17)
    ascii_font = _code_font(font_size)
    cjk_font = _code_font(font_size, cjk=True)
    cjk_bold_font = _code_font(font_size, cjk=True, bold=True)
    header = f"{source_path.relative_to(ROOT)}  |  lines {start_line}-{end_line}"
    draw.text((24, 13), header, font=header_font, fill="#FFFFFF")

    for row_index, line in enumerate(selected):
        y = header_height + 14 + row_index * line_height
        line_number = str(start_line + row_index)
        number_box = draw.textbbox((0, 0), line_number, font=line_number_font)
        number_width_px = number_box[2] - number_box[0]
        draw.text((number_width - 16 - number_width_px, y), line_number, font=line_number_font, fill="#8B96A3")
        x = number_width + horizontal_padding
        column = 0
        for character_index, character in enumerate(line):
            if character == "\t":
                next_column = ((column // 4) + 1) * 4
                x += (next_column - column) * cell_width
                column = next_column
                continue
            color, bold = styles[row_index][character_index]
            is_cjk = ord(character) > 127
            current_font = cjk_bold_font if is_cjk and bold else cjk_font if is_cjk else ascii_font
            draw.text((x, y), character, font=current_font, fill=color)
            advance = 2 if unicodedata.east_asian_width(character) in {"W", "F"} else 1
            x += advance * cell_width
            column += advance
            if x > width - horizontal_padding:
                break

    draw.rectangle((0, 0, width - 1, height - 1), outline="#BCC7D2", width=2)
    image.save(output_path)


def draw_code_architecture(path: Path) -> None:
    """Draw the implemented module flow with research boundaries."""

    width, height = 1600, 920
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _code_font(34, cjk=True, bold=True)
    heading_font = _code_font(24, cjk=True, bold=True)
    body_font = _code_font(20, cjk=True)
    draw.text((70, 42), "AUT KG Extraction Pipeline：已实现代码架构", font=title_font, fill="#17212B")

    boxes = [
        (70, 150, 360, 300, "数据适配与清洗", "IndEgo / EASG / Reviewed Gold\nprovenance + deterministic cleaning", "#E8EEF5"),
        (450, 150, 740, 300, "统一文本模块", "Raw -> UnifiedTextRecord\nevidence 必须来自原文", "#E8F3EC"),
        (830, 150, 1120, 300, "Schema + LLM 抽取", "Pydantic ontology contract\nOllama + Instructor", "#F3EDF9"),
        (1210, 150, 1500, 300, "逐层与候选关系", "entity-first / relation-second\ncandidate + binary judge", "#FFF2DB"),
        (260, 480, 600, 650, "Property Graph", "节点稳定 ID、合并、边去重\nJSON / CSV / Cypher", "#E8EEF5"),
        (690, 480, 1030, 650, "Post-processing", "别名、工具分类、方向修正\nrelation_origin + grounding", "#FDECEC"),
        (1120, 480, 1460, 650, "Evaluation", "Node/Edge P-R-F1\nontology、hallucination、stability", "#E8F3EC"),
    ]
    for x0, y0, x1, y1, title, body, fill in boxes:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=10, fill=fill, outline="#9AA9B7", width=3)
        draw.text((x0 + 20, y0 + 20), title, font=heading_font, fill="#1F4D78")
        for index, line in enumerate(body.splitlines()):
            draw.text((x0 + 20, y0 + 72 + index * 34), line, font=body_font, fill="#17212B")

    arrows = [
        ((360, 225), (450, 225)),
        ((740, 225), (830, 225)),
        ((1120, 225), (1210, 225)),
        ((1355, 300), (1290, 480)),
        ((1120, 565), (1030, 565)),
        ((690, 565), (600, 565)),
    ]
    for start, end in arrows:
        draw.line((*start, *end), fill="#2E74B5", width=5)
        direction = 1 if end[0] > start[0] else -1
        draw.polygon(
            [end, (end[0] - direction * 18, end[1] - 10), (end[0] - direction * 18, end[1] + 10)],
            fill="#2E74B5",
        )

    draw.rounded_rectangle((250, 755, 1450, 850), radius=8, fill="#F2F4F7", outline="#9AA9B7", width=2)
    draw.text((280, 775), "实验控制层：raw/unified x one-shot/layered x model x repetitions", font=heading_font, fill="#1F4D78")
    draw.text((280, 814), "统一保存 extraction、graph、validation、metrics、runtime 与错误记录", font=body_font, fill="#17212B")
    image.save(path)


def create_code_figures() -> tuple[Path, dict[str, Path]]:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    architecture = BUILD_DIR / "code_architecture.png"
    draw_code_architecture(architecture)
    specs = {
        "cleaning": ("backend/cleaning/factory_data.py", 120, 147),
        "unified": ("backend/preprocessing/unified_text.py", 110, 139),
        "llm": ("backend/llm/client.py", 43, 76),
        "schema": ("backend/schemas/egocentric_video.py", 306, 325),
        "schema_prompt": ("backend/extraction/extractor.py", 162, 193),
        "runner": ("backend/pipeline/experiment_runner.py", 171, 202),
        "layered": ("backend/pipeline/algorithm_experiment.py", 588, 607),
        "candidate": ("backend/pipeline/relation_candidate_pipeline.py", 195, 224),
        "binary_judge": ("backend/pipeline/relation_candidate_pipeline.py", 257, 290),
        "graph": ("backend/graph/property_graph.py", 81, 110),
        "postprocess": ("backend/pipeline/industrial_text_to_kg.py", 203, 237),
        "metrics": ("backend/evaluation/metrics.py", 24, 54),
        "hallucination": ("backend/evaluation/ontology_validation.py", 66, 99),
        "stability": ("backend/evaluation/stability.py", 83, 112),
        "gold": ("backend/datasets/industrial_reviewed_gold.py", 165, 198),
        "product": ("scripts/run_batch_industrial_texts_to_kg.py", 62, 95),
        "export": ("scripts/run_batch_industrial_texts_to_kg.py", 140, 169),
    }
    figures: dict[str, Path] = {}
    for key, (relative_path, start, end) in specs.items():
        source = ROOT / relative_path
        output = BUILD_DIR / f"code_{key}.png"
        draw_code_excerpt(output, source, start, end)
        figures[key] = output
    return architecture, figures


def font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def draw_grouped_bar_chart(
    path: Path,
    title: str,
    categories: list[str],
    series: list[tuple[str, list[float], str]],
    *,
    y_max: float = 1.0,
) -> None:
    width, height = 1400, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(34, True)
    axis_font = font(22)
    legend_font = font(22, True)
    value_font = font(18, True)
    draw.text((70, 35), title, font=title_font, fill="#17212B")

    left, top, right, bottom = 110, 135, 1330, 640
    for step in range(6):
        value = y_max * step / 5
        y = bottom - (bottom - top) * step / 5
        draw.line((left, y, right, y), fill="#D9E0E7", width=2)
        tick_label = f"{value:.2f}" if y_max <= 0.2 else f"{value:.1f}"
        draw.text((25, y - 12), tick_label, font=axis_font, fill="#5B6573")
    draw.line((left, top, left, bottom), fill="#5B6573", width=2)
    draw.line((left, bottom, right, bottom), fill="#5B6573", width=2)

    group_width = (right - left) / len(categories)
    bar_width = min(95, group_width / (len(series) + 1))
    for category_index, category in enumerate(categories):
        group_center = left + group_width * (category_index + 0.5)
        total_bar_width = len(series) * bar_width
        start_x = group_center - total_bar_width / 2
        for series_index, (_, values, color) in enumerate(series):
            value = values[category_index]
            x0 = start_x + series_index * bar_width
            x1 = x0 + bar_width - 10
            y0 = bottom - (bottom - top) * value / y_max
            draw.rectangle((x0, y0, x1, bottom), fill=color)
            label = f"{value:.3f}"
            bbox = draw.textbbox((0, 0), label, font=value_font)
            draw.text(((x0 + x1 - (bbox[2] - bbox[0])) / 2, y0 - 27), label, font=value_font, fill="#17212B")
        bbox = draw.textbbox((0, 0), category, font=axis_font)
        draw.text((group_center - (bbox[2] - bbox[0]) / 2, bottom + 20), category, font=axis_font, fill="#17212B")

    legend_x = 110
    legend_y = 700
    for name, _, color in series:
        draw.rectangle((legend_x, legend_y, legend_x + 24, legend_y + 24), fill=color)
        draw.text((legend_x + 34, legend_y - 2), name, font=legend_font, fill="#17212B")
        legend_x += 250
    image.save(path)


def create_charts() -> tuple[Path, Path, Path]:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    industrial = BUILD_DIR / "industrial_end_to_end_f1.png"
    relation = BUILD_DIR / "reviewed_relation_diagnostic.png"
    easg = BUILD_DIR / "easg_10scene_f1.png"

    draw_grouped_bar_chart(
        industrial,
        "Industrial 5-scene end-to-end F1",
        ["Raw+Layered", "Raw+One-shot", "Unified+Layered", "Unified+One-shot"],
        [
            ("Node F1", [0.4372, 0.5440, 0.6019, 0.5181], "#2E74B5"),
            ("Edge F1", [0.0000, 0.1421, 0.0476, 0.0857], "#C00020"),
        ],
    )
    draw_grouped_bar_chart(
        relation,
        "32-scene Gold-entity relation diagnostic",
        ["ACTS_ON", "BEFORE", "USES_TOOL"],
        [
            ("Precision", [0.8695, 0.9763, 0.5283], "#2E74B5"),
            ("Recall", [0.8222, 1.0000, 0.2979], "#3A8D5D"),
            ("F1", [0.8452, 0.9880, 0.3810], "#C00020"),
        ],
    )
    draw_grouped_bar_chart(
        easg,
        "EASG 10-scene real-LLM F1",
        ["Raw+Layered", "Raw+One-shot", "Unified+Layered", "Unified+One-shot"],
        [
            ("Node F1", [0.8898, 0.7976, 0.8685, 0.5498], "#2E74B5"),
            ("Edge F1", [0.4086, 0.4086, 0.3755, 0.4086], "#C00020"),
        ],
    )
    return industrial, relation, easg


def add_experiment_descriptor(
    doc: Document,
    *,
    purpose: str,
    implementation: str,
    metrics: str,
    result: str,
    problem: str,
) -> None:
    add_labeled(doc, "实验目的", purpose)
    add_labeled(doc, "实现方式", implementation)
    add_labeled(doc, "评价标准", metrics)
    add_labeled(doc, "主要成果", result)
    add_labeled(doc, "问题与边界", problem)


def add_module_details(
    doc: Document,
    *,
    role: str,
    input_output: str,
    implementation: str,
    experiment: str,
    references: str,
    boundary: str,
) -> None:
    """Write the explanation that belongs directly below a code figure."""

    add_labeled(doc, "模块作用", role)
    add_labeled(doc, "输入与输出", input_output)
    add_labeled(doc, "实现方式与关键逻辑", implementation)
    add_labeled(doc, "对应实验结论", experiment)
    add_labeled(doc, "参考来源", references)
    add_labeled(doc, "当前边界", boundary)


def classify_result_dir(name: str) -> tuple[str, str]:
    primary = {
        "local_experiment_2026-05-26_r1_complete",
        "real_llm_long_2026-06-19",
        "hospital_postprocess_long_2026-06-19",
        "hospital_text_processing_long_2026-06-19",
        "easg_module_experiment_full_2026-06-20",
        "easg_real_llm_lightweight_realtest_10scene_2026-06-20",
        "industrial_real_llm_lightweight_final_2026-06-20",
        "industrial_full_after_relation_fix_2026-07-05",
        "industrial_full_after_relation_fix_rerun2_2026-07-05",
        "industrial_candidate_pipeline_eval_2026-07-05",
        "relation_candidate_ablation_2026-07-07",
        "industrial_reviewed_gold_v2_relation_diagnostic_2026-08-01",
        "industrial_reviewed_gold_combined_relation_diagnostic_2026-08-01",
    }
    incomplete = {
        "easg_real_llm_small_2026-06-20",
        "easg_real_llm_smoke_2026-06-20",
        "industrial_reviewed_gold_v2_pilot_5scene_2026-08-01",
        "industrial_reviewed_gold_v2_pilot_5scene_tokens1800_2026-08-01",
        "integrated_real_llm_all_codex_attempt_2026-06-19",
    }
    if name in primary:
        return "主要证据", "进入正文结论"
    if name in incomplete:
        return "未完成/失败运行", "只记录问题，不汇总为结论"
    if "smoke" in name or "quick" in name or "stability" in name:
        return "Smoke/支持性验证", "用于工程可行性或扩大实验前检查"
    if name.startswith(("real_candidate", "staged_", "diagnose_", "real_minimal", "real_egocentric")):
        return "开发诊断", "用于定位代码问题，不作为研究结果"
    return "支持性运行", "作为版本或回归证据"


def build_document() -> Path:
    industrial_chart, relation_chart, easg_chart = create_charts()
    architecture_chart, code_figures = create_code_figures()
    doc = Document()
    configure_document(doc)

    # First-page memo masthead.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run("TU BERLIN · INDUSTRIAL AUTOMATION PROJECT")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("AUT KG Extraction Pipeline\n实验与代码全量总结报告")
    set_run_font(title_run, size=24, color=INK, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle_run = subtitle.add_run("代码架构、关键实现、实验结果、问题与改进措施")
    set_run_font(subtitle_run, size=13, color=MUTED)

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["报告日期", "2026-08-03"],
            ["覆盖期间", "2026-05-26 至 2026-08-01"],
            ["主要模型", "本地 Ollama llama3.1:8b；对照模型 qwen2.5:7b、llama3.2:latest"],
            ["最终研究对象", "视频衍生工业操作文本到程序性知识图谱"],
            ["Gold 原则", "仅允许 source-supported facts，不加入无证据推理事实"],
        ],
        [1700, 7660],
        font_size=9,
        first_column_bold=True,
    )
    rule = doc.add_paragraph()
    set_paragraph_border_bottom(rule, color=BLUE, size=12)

    doc.add_page_break()
    doc.add_heading("执行摘要", level=1)
    add_callout(
        doc,
        "当前最重要结论",
        "系统已经形成可运行的工业文本到知识图谱链路，节点抽取与格式稳定性已有明确进展；关系抽取仍是主要瓶颈。最新 32 场景 Gold-entity 关系诊断达到 Precision 0.8926、Recall 0.8333、F1 0.8620，但这不是端到端 LLM 分数。工业端到端实验的最佳 Node F1 为 0.6019，最佳 Edge F1 为 0.1421。",
    )
    add_body(
        doc,
        "本报告汇总项目中已经执行的医院模块筛选、EASG 模拟与真实 LLM 实验、工业 5 场景端到端实验、关系候选消融，以及人工审核 V2+V3 Gold 的关系诊断。报告不把 smoke test、模拟实验、Gold-entity oracle 诊断和端到端 LLM 实验混为同一种证据。",
    )
    add_bullets(
        doc,
        [
            "统一文本结构在多轮实验中主要改善 schema/JSON 可用性、ID 显式性和部分节点质量，但不保证 Edge F1 必然提高。",
            "逐层抽取比一次性全图抽取更容易运行，但当前 layered relation prompt 并未稳定超过 one-shot 的关系质量。",
            "节点标准化、关系别名映射、方向修正、重复边合并和 hallucination 检查适合作为产品后处理，但 raw 与 normalized 指标必须分开报告。",
            "LLM binary judge 可以过滤部分错误边，却不能恢复候选生成阶段漏掉的关系；候选 recall 是后续判断器的上限。",
            "USES_TOOL 是最新 32 场景诊断中最弱的关系，Recall 仅 0.2979。",
        ],
    )

    doc.add_heading("1. 报告范围与证据等级", level=1)
    add_body(
        doc,
        "项目研究范围从 video-to-text conversion 之后开始，不训练原始视频模型。实验输入包括工业 scene description、视频 narration 或 annotation-derived text。医院数据只用于筛选模块；EASG 用于公开人工 scene-graph 支持性验证；最终研究结论应以工业 reviewed Gold 为主。",
    )
    add_table(
        doc,
        ["证据等级", "数据/实验", "允许支持的结论", "不能支持的结论"],
        [
            ["A：主要工业证据", "工业 5 场景端到端；32 场景 reviewed-Gold 关系诊断", "当前工业节点/关系性能与模块瓶颈", "大规模泛化能力"],
            ["B：支持性公开证据", "EASG annotation-only 实验", "抽取流程、JSON 稳定性与 scene-graph 适配", "工业制造领域最终性能"],
            ["C：模块筛选", "医院 Gold pilot、stress、post-processing long", "prompt、chunking、normalization 等模块价值", "医学 KG 或工业 KG 最终结果"],
            ["D：工程验证", "smoke、staged、diagnostic、partial runs", "代码可运行性与失败原因", "研究结论或模型排名"],
        ],
        [1450, 2380, 2860, 2670],
        font_size=8.2,
        first_column_bold=True,
    )
    add_source_note(doc, "项目内部实验报告与 results/ 目录；证据等级由本报告按研究边界整理。")

    doc.add_heading("1.1 数据集角色", level=2)
    add_table(
        doc,
        ["数据集", "规模", "用途", "限制"],
        [
            ["Hospital Gold Case", "单 admission，47 节点/48 边后处理基准", "筛选 schema、分层抽取、分块、模型、后处理", "不是最终工业数据"],
            ["EASG", "6250 annotation scenes；真实 LLM 子集 1/3/10 scenes", "公开人工 scene-graph 与 annotation-derived text", "非制造业专用；不是自然 transcript"],
            ["Industrial seed benchmark", "5 scenes", "早期和 2026-07 工业端到端 LLM", "样本量小"],
            ["Reviewed Gold V2", "12 scenes，272 实体，325 关系", "人工审核后关系诊断与 smoke", "仍是开发集"],
            ["Reviewed Gold V3", "20 scenes，402 实体，503 关系", "扩展动作、工具和关系覆盖", "建议抽样二审"],
            ["Combined V2+V3", "32 scenes，674 实体，828 关系，30 video IDs", "当前 WP7 正式 pilot scope", "尚未完成全量重复端到端 LLM"],
        ],
        [1800, 1800, 3100, 2660],
        font_size=8.2,
        first_column_bold=True,
    )

    doc.add_heading("1.2 Gold Graph 边界", level=2)
    add_bullets(
        doc,
        [
            "Gold entity 和 Gold relation 必须有 source_text 或 evidence span 支持。",
            "人工推理、医学常识补充和 AI 生成但原文未表达的关系不进入 Gold。",
            "候选标注只有经过人工审核并标记 accepted 后才可以作为正式 pilot Gold。",
            "自动后处理修正必须记录 relation_origin；不能把 postprocessed 结果描述为 LLM 原始能力。",
        ],
    )

    doc.add_heading("2. 评价标准", level=1)
    add_body(doc, "评价同时覆盖正确性、完整性、幻觉、schema 合规性、稳定性和工程可运行性。节点与关系分开计算，避免实体分数掩盖关系错误。")
    add_table(
        doc,
        ["研究维度", "指标", "计算/解释"],
        [
            ["Correctness 正确性", "Precision", "TP / (TP + FP)：预测事实中有多少与 Gold 匹配"],
            ["Completeness 完整性", "Recall", "TP / (TP + FN)：Gold 事实中有多少被覆盖"],
            ["综合质量", "F1", "2 × Precision × Recall / (Precision + Recall)"],
            ["节点质量", "Node P/R/F1", "按实体类型与规范化名称/ID 匹配 Action、Tool、Object 等"],
            ["关系质量", "Edge P/R/F1", "按 relation type、source endpoint、target endpoint 匹配"],
            ["格式可用性", "Schema success / Error count", "Pydantic/JSON 是否可解析；超时与解析错误是否发生"],
            ["本体约束", "Ontology conformance", "关系类型、domain/range 和端点类型是否符合 schema"],
            ["幻觉", "Unsupported node/edge count", "输出事实是否缺少原文、ontology 或 source CSV 支持"],
            ["稳定性", "Graph overlap / relation agreement / F1 std", "多次运行图集合的 Jaccard 重合与指标方差"],
            ["工程成本", "Runtime", "单次或条件平均推理时间；用于判断本地模型可行性"],
        ],
        [1900, 2050, 5410],
        font_size=8.5,
        first_column_bold=True,
    )
    add_callout(
        doc,
        "解释规则",
        "Graph overlap = 1.0 只表示模型重复输出一致，不表示输出正确。Gold-entity diagnostic 提供正确实体端点，因此只评价关系阶段上限；其 F1 不能当作端到端 LLM F1。",
        risk=True,
    )
    doc.add_heading("2.1 Raw 与 Normalized 双指标", level=2)
    add_body(doc, "Raw metrics 反映模型原始输出能力；Normalized metrics 反映节点标准化、别名映射、方向修正和重复合并后的产品结果。两者必须同时保留，防止后处理掩盖上游弱点。")

    doc.add_page_break()
    doc.add_heading("3. 实验总览", level=1)
    overview_rows = [
        ["E01", "早期工业 Raw vs Unified", "真实 LLM / 5 scenes / 10 calls", "验证链路与 schema 稳定性"],
        ["E02", "医院 P1-P6 Pilot", "真实 LLM / 6 个变量组", "筛选 prompt、layering、chunk、model、stability"],
        ["E03", "医院 Stress", "真实 LLM / prompt、chunk、noise", "让过于简单的任务产生区分度"],
        ["E04", "医院后处理 Long", "确定性 / 7200 runs", "验证 normalization、merge、hallucination"],
        ["E05", "EASG 模拟 Quick/Full", "模拟 / 600 + 125000 runs", "验证实验框架与模块假设"],
        ["E06", "EASG smoke + JSON repair", "真实 LLM / 1 scene", "定位输出格式失败"],
        ["E07", "EASG 3-scene", "真实 LLM / 24 runs", "初步扩大 raw/unified × strategy"],
        ["E08", "EASG 10-scene", "真实 LLM / 80 runs", "验证 parser repair 与条件差异"],
        ["E09", "工业 5-scene baseline", "真实 LLM / 40 runs", "最终工业方向的早期端到端基线"],
        ["E10", "关系解析修复 + 重跑", "真实 LLM / 40 + 40 runs", "修复关系 payload 和 label normalization"],
        ["E11", "候选关系架构评估", "真实 LLM / 40 runs", "验证 deterministic candidate fallback"],
        ["E12", "候选关系消融", "真实 LLM / 30 method-runs", "比较 minimal、rules、LLM judge"],
        ["E13", "Reviewed Gold V2 smoke/诊断", "真实 LLM smoke + 12-scene oracle diagnostic", "检验长 transcript 与关系覆盖"],
        ["E14", "Combined V2+V3 诊断", "32-scene Gold-entity diagnostic", "最新关系模块覆盖与瓶颈"],
    ]
    add_table(doc, ["ID", "实验", "实现规模", "主要目的"], overview_rows, [700, 2400, 2600, 3660], font_size=7.9, first_column_bold=True)

    doc.add_page_break()
    doc.add_heading("4. 早期工业实验", level=1)
    doc.add_heading("E01 早期本地工业 Raw vs Unified 试运行", level=2)
    add_experiment_descriptor(
        doc,
        purpose="确认 5 场景工业 benchmark 能否完整经过本地 llama3.1:8b、Pydantic schema、property graph 和 evaluation；初步比较 raw 与 unified。",
        implementation="每个场景在 raw 和 unified 条件下各运行 1 次，共 10 次模型调用；仍是单阶段早期抽取合同。",
        metrics="Schema 成功数、Node F1、Relation F1、schema 合规率。",
        result="Raw 为 3/5 schema 成功，Node F1 0.4638、Relation F1 0.1818；Unified 为 5/5 成功，Node F1 0.5000、Relation F1 0.0000。",
        problem="每条件仅运行 1 次；unified 提高了结构可用性，却没有改善关系。该实验促使项目把 Node 与 Edge 分开评价。",
    )
    add_table(
        doc,
        ["输入", "运行", "Schema 成功", "Node F1", "Relation F1", "合规率"],
        [["Raw", 5, 3, "0.4638", "0.1818", "0.6000"], ["Unified", 5, 5, "0.5000", "0.0000", "1.0000"]],
        [1400, 1100, 1600, 1600, 1760, 1900],
        font_size=8.5,
        first_column_bold=True,
    )
    add_source_note(doc, "docs/results/local_experiment_2026-05-26_r1.md")

    doc.add_page_break()
    doc.add_heading("5. 医院模块筛选实验", level=1)
    add_callout(doc, "定位", "医院 Gold Case 用来回答“哪些模块值得保留”，不用于医院知识图谱结论，也不替代最终工业实验。")

    doc.add_heading("E02 医院 P1-P6 Pilot", level=2)
    add_experiment_descriptor(
        doc,
        purpose="以小型可控 Gold graph 快速验证模型、schema、分层抽取、文本分块、模型差异与重复稳定性。",
        implementation="使用 hospital admission 29366372 的 source CSV 与人工 Gold；主要模型为 llama3.1:8b，并加入 qwen2.5:7b、llama3.2:latest 对照。",
        metrics="Node/Edge P/R/F1、Extra nodes/edges、timeout、graph overlap、F1 标准差、runtime。",
        result="严格 schema 在最小任务达到 Node/Edge F1 1.0；早期 one-shot 全图超时，layered 能返回结果；模型对关系方向表现不同；稳定性可能稳定地重复错误。",
        problem="原始 P1/P3/P4 任务过于简单，多个条件满分，无法区分模块价值，因此追加 stress 与 long rerun。",
    )
    add_table(
        doc,
        ["子实验", "变量", "早期主要结果", "实验意义"],
        [
            ["P1", "strict schema", "Node/Edge F1 = 1.0/1.0", "验证调用与评价链路"],
            ["P2", "one-shot vs layered", "one-shot 180s timeout；layered 0.6190/0.5647", "逐层拆任务更可运行"],
            ["P3", "chunk 12/6/3/1", "简单 LabEvent 全部 1.0/1.0", "任务过易，需 stress"],
            ["P4", "raw vs unified", "简单文本均 1.0/1.0", "不能证明 unified 提升"],
            ["P5", "3 models", "llama3.2 Node 0.9231、Edge 0；其余满分", "模型会影响关系方向"],
            ["P6", "3 次重复", "llama3.2 overlap 1.0 但 Edge F1 0", "稳定性不等于正确性"],
        ],
        [900, 2000, 3180, 3280],
        font_size=8.0,
        first_column_bold=True,
    )
    add_source_note(doc, "docs/results/hospital_gold_pilot_2026-06-04.md；docs/results/hospital_gold_pilot_figures_zh.md")

    doc.add_heading("E03 医院 P1/P3/P4 Stress 与 Long Rerun", level=2)
    add_experiment_descriptor(
        doc,
        purpose="增加 prompt 松散度、输入噪声和密集 narrative，使变量差异足够明显。",
        implementation="P1 比较 loose/strict prompt；P3 在多实体多关系 narrative 上改变 chunk size；P4 比较无 canonical ID 的 raw noisy text 与字段化 unified input。后续 real_llm_long 使用统一 runner 重跑。",
        metrics="Raw/normalized Node F1、Edge F1、FP/FN、timeout、错误类型。",
        result="Strict prompt 稳定达到 1.0/1.0；long run 中 loose raw 为 0.7143/0.6667。P4 long run 的 raw noisy 为 0.6250/0.6000，unified 为 1.0/1.0。大块 dense narrative 曾出现 120s timeout。",
        problem="不同 stress 版本的噪声强度和后处理版本不同，数值不可直接当作同一条件重复。小分块恢复运行，但节点 ID/label mismatch 增加。",
    )
    add_table(
        doc,
        ["实验", "条件", "Raw Node F1", "Raw Edge F1", "Normalized Edge F1", "问题"],
        [
            ["P1 long", "Loose", "0.7143", "0.6667", "0.8000", "非标准/额外关系"],
            ["P1 long", "Strict", "1.0000", "1.0000", "1.0000", "无"],
            ["P4 long", "Raw noisy", "0.6250", "0.6000", "0.7500", "ID 与方向不稳定"],
            ["P4 long", "Unified", "1.0000", "1.0000", "1.0000", "受控结构"],
        ],
        [1300, 1450, 1500, 1500, 1750, 1860],
        font_size=8.2,
        first_column_bold=True,
    )
    add_source_note(doc, "docs/results/hospital_gold_pilot_p134_stress_zh.md；results/real_llm_long_2026-06-19/summary.md")

    doc.add_page_break()
    doc.add_heading("E04 医院 Post-processing Long Benchmark", level=2)
    add_experiment_descriptor(
        doc,
        purpose="压力测试 deterministic 后处理，不把 LLM 随机性混入节点标准化、关系合并、方向修正和 hallucination 检查。",
        implementation="47 Gold nodes、48 Gold edges；6 种噪声条件，每条件 1200 次，共 7200 次；seed 20260619。",
        metrics="Raw/normalized Node/Edge F1、hallucinated edges、wrong direction、merged duplicates、平均耗时。",
        result="alias_noise 从 raw Node/Edge F1 0/0 恢复到 normalized 1/1；reverse_duplicate_noise 修正 48 个反向关系并合并 48 个重复边；severe_invalid 仍只有 normalized Node 0.2909、Edge 0.1825。",
        problem="后处理只能修正表示、方向和重复问题，无法恢复上游漏抽事实；normalized 高分不能代替 raw 模型能力。",
    )
    add_table(
        doc,
        ["Condition", "Raw Node", "Norm Node", "Raw Edge", "Norm Edge"],
        [
            ["clean_canonical", "1.0000", "1.0000", "1.0000", "1.0000"],
            ["alias_noise", "0.0000", "1.0000", "0.0000", "1.0000"],
            ["reverse_duplicate_noise", "0.0000", "1.0000", "0.0000", "1.0000"],
            ["hallucination_stress", "0.0000", "0.9495", "0.0000", "1.0000"],
            ["mixed_noisy_reduced", "0.2106", "0.8471", "0.1443", "0.7878"],
            ["severe_invalid", "0.0000", "0.2909", "0.0000", "0.1825"],
        ],
        [2600, 1690, 1690, 1690, 1690],
        font_size=8.4,
        first_column_bold=True,
    )
    add_source_note(doc, "results/hospital_text_processing_long_2026-06-19/summary.md；postprocess_long_summary.csv")

    doc.add_page_break()
    doc.add_heading("6. EASG 实验", level=1)
    doc.add_heading("E05 EASG 模块模拟 Quick 与 Full", level=2)
    add_experiment_descriptor(
        doc,
        purpose="在不调用真实 LLM 的情况下验证 raw/unified 与 one-shot/layered 的实验设计、稳定性计算和 hallucination 统计。",
        implementation="Quick: 50 scenes × 3 repetitions × 4 conditions = 600 runs；Full: 6250 scenes × 5 repetitions × 4 conditions = 125000 runs。输出通过受控随机扰动模拟。",
        metrics="Node/Edge F1、graph overlap、hallucinated nodes/edges。",
        result="Full simulation 中 unified+layered 最佳：Node F1 0.8747、Edge F1 0.7423、overlap 0.8279，hallucinated edges 0.0011。",
        problem="模拟结果只能验证实验框架和模块假设，不能写成 llama3.1:8b 的真实性能。",
    )
    add_table(
        doc,
        ["Condition", "Strategy", "Node F1", "Edge F1", "Overlap", "Hallucinated N/E"],
        [
            ["Raw", "One-shot", "0.7986", "0.4593", "0.5670", "0.2030 / 0.0249"],
            ["Raw", "Layered", "0.8328", "0.5880", "0.6571", "0.1195 / 0.0088"],
            ["Unified", "One-shot", "0.8432", "0.5905", "0.6751", "0.1025 / 0.0080"],
            ["Unified", "Layered", "0.8747", "0.7423", "0.8279", "0.0405 / 0.0011"],
        ],
        [1500, 1500, 1450, 1450, 1350, 2110],
        font_size=8.2,
        first_column_bold=True,
    )
    add_source_note(doc, "results/easg_module_experiment_quick_2026-06-20；results/easg_module_experiment_full_2026-06-20")

    doc.add_heading("E06-E08 EASG 真实 LLM 扩大实验", level=2)
    add_experiment_descriptor(
        doc,
        purpose="逐步从 1 scene smoke 扩展到 3 scene 和 10 scene，验证 JSON repair、layered extraction 与 unified input 在真实本地模型上的作用。",
        implementation="使用 llama3.1:8b 轻量 JSON contract。Parser 只修复 code fence、前后缀、comment、尾逗号、未加引号 key、单引号和安全 top-level list，不补写图事实。",
        metrics="Node/Edge F1、error count、graph overlap、runtime；10 scenes × 2 repetitions × 4 conditions = 80 runs。",
        result="加入 JSON repair 后，10-scene 四个条件均 0 error；最佳 Node F1 为 raw+layered 0.8898；Edge F1 约 0.3755-0.4086。",
        problem="EASG 输入是 annotation-derived text，不是自然视频 transcript，也不是制造业专用数据；overlap 1.0 不代表正确。",
    )
    add_figure(doc, easg_chart, "图 1  EASG 10 场景真实 LLM 的 Node/Edge F1", "EASG ten-scene real LLM F1 grouped bar chart")
    add_table(
        doc,
        ["Condition", "Strategy", "Runs", "Errors", "Node F1", "Edge F1", "Runtime(s)"],
        [
            ["Raw", "Layered", 20, 0, "0.8898", "0.4086", "26.08"],
            ["Raw", "One-shot", 20, 0, "0.7976", "0.4086", "29.57"],
            ["Unified", "Layered", 20, 0, "0.8685", "0.3755", "25.46"],
            ["Unified", "One-shot", 20, 0, "0.5498", "0.4086", "28.06"],
        ],
        [1300, 1500, 900, 900, 1350, 1350, 2060],
        font_size=8.2,
        first_column_bold=True,
    )
    add_source_note(doc, "docs/results/easg_real_llm_lightweight_realtest_10scene_2026-06-20.md")

    doc.add_page_break()
    doc.add_heading("7. 工业端到端真实 LLM 实验", level=1)
    doc.add_heading("E09 2026-06-20 工业 5 场景基线", level=2)
    add_experiment_descriptor(
        doc,
        purpose="在最终工业方向上比较 raw/unified × one-shot/layered，并检查真实本地模型格式稳定性。",
        implementation="5 scenes × 2 repetitions × 4 conditions = 40 runs；llama3.1:8b；轻量 JSON contract。",
        metrics="Node/Edge F1、errors、graph overlap、runtime。",
        result="当时最佳 Node F1 为 unified+one-shot 0.5181；Edge F1 仅 0-0.0250；两个 unified 条件 0 error，raw+one-shot 有 2 error。",
        problem="关系 payload 解析和 relation label normalization 尚有 bug，低估了显式 LLM 关系；工业文本明显比 EASG 困难。",
    )

    doc.add_heading("E10 关系解析修复与全量重跑", level=2)
    add_experiment_descriptor(
        doc,
        purpose="修复 grouped relation payload 和 ACTS_ON/USES_TOOL 等 schema label 被错误清洗的问题，并验证结果能否复现。",
        implementation="新增 grouped relation shape normalization、normalize_relation_type() 和 relation_origin 统计；同样运行 40 次并再重复一轮。",
        metrics="Node/Edge F1、errors、graph overlap、relation origin counts、两次全量指标 delta。",
        result="最佳 Node F1 更新为 unified+layered 0.6019；最佳 Edge F1 为 raw+one-shot 0.1421；第二次全量重跑所有质量指标 delta 均为 0。32/40 runs 至少包含一条 llm_extracted edge。",
        problem="Layered 提高了节点，但 Edge F1 低于 raw+one-shot；大量关系仍依赖 evidence_fallback、sequence_fallback 或 postprocessed。",
    )
    add_figure(doc, industrial_chart, "图 2  修复后的工业 5 场景端到端 Node/Edge F1", "Industrial five-scene end-to-end Node and Edge F1 grouped bar chart")
    add_table(
        doc,
        ["Condition", "Strategy", "Runs", "Errors", "Node F1", "Edge F1", "Runtime(s)"],
        [
            ["Raw", "Layered", 10, 0, "0.4372", "0.0000", "35.67"],
            ["Raw", "One-shot", 10, 0, "0.5440", "0.1421", "39.72"],
            ["Unified", "Layered", 10, 0, "0.6019", "0.0476", "52.85"],
            ["Unified", "One-shot", 10, 0, "0.5181", "0.0857", "39.79"],
        ],
        [1300, 1500, 900, 900, 1350, 1350, 2060],
        font_size=8.2,
        first_column_bold=True,
    )
    add_table(
        doc,
        ["Relation origin", "Count", "解释"],
        [
            ["llm_extracted", 94, "模型显式关系"],
            ["evidence_fallback", 46, "基于文本证据的保守补全"],
            ["sequence_fallback", 44, "按动作顺序生成 BEFORE"],
            ["postprocessed", 116, "例如 OBSERVED_IN 等后处理边"],
        ],
        [2400, 1300, 5660],
        font_size=8.5,
        first_column_bold=True,
    )
    add_source_note(doc, "docs/results/industrial_full_after_relation_fix_2026-07-05.md；industrial_full_after_relation_fix_rerun2_2026-07-05.md")

    doc.add_page_break()
    doc.add_heading("8. 候选关系方法实验", level=1)
    doc.add_heading("E11 Generic Relation Candidate Pipeline", level=2)
    add_experiment_descriptor(
        doc,
        purpose="把自由生成关系改造成通用接口 RelationCandidateGenerator → RelationCandidateScorer → RelationValidationReport。",
        implementation="从已抽取实体和 source evidence 生成 schema-valid candidate；deterministic scorer 作为保守 fallback；仍运行 5 scenes × 2 repetitions × 4 conditions。",
        metrics="Node/Edge F1、error、平均候选/接受边数、relation origin。",
        result="架构成功集成且没有明显边爆炸，但 deterministic candidate fallback 没有提高 Edge F1。统一 layered 仍为 Node 0.6019、Edge 0.0476。",
        problem="Deterministic scorer 接受的新增候选多为错误端点或不匹配 Gold；raw one-shot 有 2 次 API timeout，因此 Node F1 暂降。",
    )
    add_source_note(doc, "docs/results/industrial_candidate_pipeline_eval_2026-07-05.md")

    doc.add_page_break()
    doc.add_heading("E12 Minimal vs Deterministic vs LLM Binary Judge 消融", level=2)
    add_experiment_descriptor(
        doc,
        purpose="直接验证候选关系新方向能否提高 Edge F1，并判断 LLM binary judge 是提高 precision 还是恢复 recall。",
        implementation="统一输入；5 scenes × 2 repetitions；同一次实体抽取分别进入 current minimal、deterministic rules、LLM binary judge。Judge 只能回答 candidate_id、is_supported、confidence、evidence_text、reason。",
        metrics="Node/Edge F1、edge agreement、predicted/candidate/accepted edges、runtime。",
        result="Node F1 三种方法均 0.6019；Edge F1 为 0.0476、0.0476、0.0564。LLM judge 将平均预测边从 8.6 降到 5.8，但运行时间增至 81.84s。",
        problem="改进幅度很小；judge 是 precision filter，不能补回 generator 未生成的关系。下一步必须先提高 candidate recall。",
    )
    add_table(
        doc,
        ["Method", "Runs", "Errors", "Node F1", "Edge F1", "Pred edges", "Runtime(s)"],
        [
            ["current_minimal_candidate", 10, 0, "0.6019", "0.0476", "7.60", "19.34"],
            ["candidate_deterministic_rules", 10, 0, "0.6019", "0.0476", "8.60", "0.00"],
            ["candidate_llm_binary_judge", 10, 0, "0.6019", "0.0564", "5.80", "81.84"],
        ],
        [2600, 850, 850, 1300, 1300, 1300, 1160],
        font_size=8.0,
        first_column_bold=True,
    )
    add_source_note(doc, "docs/results/relation_candidate_ablation_2026-07-07.md")

    doc.add_page_break()
    doc.add_heading("9. 人工审核 Gold 实验", level=1)
    doc.add_heading("E13 Reviewed Gold V2 Smoke 与关系诊断", level=2)
    add_experiment_descriptor(
        doc,
        purpose="把人工审核后的 12 场景 Gold 接入统一 benchmark adapter，并检查真实长 transcript 与关系候选覆盖。",
        implementation="先运行单场景 lightweight smoke；随后把 Gold entities 直接提供给 relation candidate generator，执行 12-scene relation-stage diagnostic。",
        metrics="Smoke 的 Node/Edge F1、errors、runtime；诊断的 candidate coverage recall、accepted P/R/F1、按 relation type 统计。",
        result="初始 token budget 下 4 条 smoke 全失败；提高到 3000 tokens 后单场景 unified+layered 达到 Node 0.9655、Edge 0.8657，但耗时 152.88s。12-scene 关系诊断 P/R/F1 为 0.8900/0.8462/0.8675。",
        problem="单场景高分不能泛化；5-scene pilot 只产生 2 条 checkpoint，没有完整汇总。自动 summary 仍残留 EASG 通用措辞，属于报告模板 bug。",
    )
    add_table(
        doc,
        ["Relation", "Gold", "Candidates", "Precision", "Recall", "F1"],
        [
            ["ACTS_ON", 170, 162, "0.8210", "0.7824", "0.8012"],
            ["BEFORE", 133, 133, "1.0000", "1.0000", "1.0000"],
            ["USES_TOOL", 22, 14, "0.6429", "0.4091", "0.5000"],
        ],
        [1800, 1150, 1500, 1600, 1600, 1710],
        font_size=8.4,
        first_column_bold=True,
    )
    add_source_note(doc, "docs/results/industrial_reviewed_gold_v2_relation_diagnostic_2026-08-01.md；results/industrial_reviewed_gold_v2_smoke_tokens3000_2026-08-01/summary.md")

    doc.add_heading("E14 Combined V2+V3 32 场景关系诊断", level=2)
    add_experiment_descriptor(
        doc,
        purpose="在更完整的人工关系 Gold 上重新估计 candidate coverage 上限，并定位具体关系类型瓶颈。",
        implementation="合并 V2 12 scenes 与 V3 20 scenes，保留原文件；共 32 scenes、30 video IDs、674 entities、828 relations。使用 Gold entities 运行 deterministic candidate diagnostic。",
        metrics="Candidate coverage recall、deterministic accepted P/R/F1、每类关系 P/R/F1、missed Gold edges。",
        result="总体 Precision 0.8926、Recall 0.8333、F1 0.8620；BEFORE F1 0.9880，ACTS_ON 0.8452，USES_TOOL 0.3810。共漏掉 138 条人工关系。",
        problem="USES_TOOL Recall 只有 0.2979；当前 generator 未覆盖多动作共享工具、跨短句工具证据和部分 action-tool endpoint 对齐。该实验不是端到端 LLM 分数。",
    )
    add_figure(doc, relation_chart, "图 3  32 场景人工 Gold 的关系类型诊断", "Thirty-two-scene reviewed Gold relation precision recall and F1 chart")
    add_table(
        doc,
        ["Relation", "Gold", "Candidates", "Matched", "Precision", "Recall", "F1"],
        [
            ["ACTS_ON", 405, 383, 333, "0.8695", "0.8222", "0.8452"],
            ["BEFORE", 329, 337, 329, "0.9763", "1.0000", "0.9880"],
            ["USES_TOOL", 94, 53, 28, "0.5283", "0.2979", "0.3810"],
        ],
        [1700, 1000, 1300, 1100, 1420, 1420, 1420],
        font_size=8.2,
        first_column_bold=True,
    )
    add_source_note(doc, "docs/results/industrial_reviewed_gold_combined_relation_diagnostic_2026-08-01.md")

    doc.add_page_break()
    doc.add_heading("10. 代码架构与模块实现", level=1)
    add_body(
        doc,
        "本章依据当前仓库中的真实代码说明系统如何从工业文本生成知识图谱。所有代码图片均由报告脚本直接读取仓库源码并按真实行号生成，不是重新排版的伪代码。每张代码图下方依次说明模块作用、输入输出、实现逻辑、实验结论、参考来源与实现边界。",
    )
    add_callout(
        doc,
        "代码来源边界",
        "项目使用 Pydantic、Instructor、OpenAI Python client 和 Ollama 的公开 API；本体约束、Property Graph、关系候选与评估设计参考相关论文或开源项目。数据清洗规则、evidence 校验、实验编排、关系来源标记、工业后处理和本项目 schema 适配均为本仓库实现，没有整段复制外部项目代码。",
    )

    doc.add_heading("10.1 总体代码架构", level=2)
    add_figure(
        doc,
        architecture_chart,
        "图 4  已实现的代码模块与数据流",
        "Implemented architecture of the AUT KG extraction pipeline",
    )
    add_body(
        doc,
        "系统由两条主链组成。第一条是 Data Cleaning Module：数据适配、确定性清洗、Raw/Unified 表示和证据保留。第二条是 Text-to-KG Module：schema 约束抽取、逐层抽取、关系候选、Property Graph、后处理与评估。实验控制层对模型、输入形式、抽取策略和重复次数进行组合，并保存每次运行的原始响应、结构化结果、图、验证报告、指标和耗时。",
    )
    add_table(
        doc,
        ["模块", "主要文件", "核心职责", "实验中的作用"],
        [
            ["数据适配/清洗", "backend/cleaning/; backend/datasets/", "来源解析、字段清理、provenance", "比较 raw 与统一结构的前提"],
            ["统一文本", "backend/preprocessing/unified_text.py", "事实分栏、evidence 约束", "提高 schema/JSON 可用性"],
            ["Schema/LLM", "backend/schemas/; backend/llm/; backend/extraction/", "Pydantic 合同与本地结构化抽取", "限制实体和关系输出空间"],
            ["逐层抽取", "backend/pipeline/algorithm_experiment.py", "实体优先、关系后置", "one-shot 对照条件"],
            ["关系候选", "backend/pipeline/relation_candidate_pipeline.py", "候选生成、评分、验证报告", "关系消融和 recall 诊断"],
            ["图与后处理", "backend/graph/; industrial_text_to_kg.py", "合并、规范化、方向和来源", "产品 V1 稳定输出"],
            ["评估", "backend/evaluation/", "P/R/F1、hallucination、stability", "回答 correctness/completeness 等问题"],
            ["一键入口", "scripts/run_*industrial_text*_kg.py", "单文件/批量输出和导出", "交付与演示入口"],
        ],
        [1600, 2600, 2650, 2510],
        font_size=7.5,
        first_column_bold=True,
    )

    doc.add_page_break()
    doc.add_heading("10.2 Exposé 原计划与当前代码方法对照", level=2)
    add_callout(
        doc,
        "结论先行",
        "Exposé 中的完整方法仍是研究蓝图；当前代码实现的是可在现有工业 scene-level 文本和本地小模型上运行、评估和审计的阶段性版本。当前 entity-first/relation-second 逐层抽取不等于完整的 dependency-graph/tree decomposition，当前 Property Graph 去重也不等于独立局部子图的全局冲突合并。",
    )

    doc.add_heading("10.2.1 Exposé 中原计划采用的方法", level=3)
    add_body(
        doc,
        "开题报告提出的是 ACONIC-inspired decomposition pipeline：先把较长的教学视频转录文本显式建模为程序依赖结构，再沿依赖结构分解，最后分别抽取和合并。其完整流程如下。",
    )
    add_numbered(
        doc,
        [
            "Preprocessing and segmentation：清理并切分 instructional video transcripts，保留动作、时间、条件和资源线索。",
            "Procedural dependency graph construction（WP3）：构造表示 Action、temporal relation、condition 与 shared resource 的程序依赖图。",
            "Locally coherent subtask decomposition（WP4）：对依赖图进行 tree decomposition，将长流程划分为局部一致的程序子任务。",
            "Local ontology-guided extraction：针对每个子任务使用本体定义、任务说明和 in-context examples，独立生成 PKO-oriented triples/local subgraphs。",
            "Global subgraph merge（WP5）：使用同类型约束、字符串相似度、置信度与 provenance 统一重复实体、解决属性冲突，并合并为全局程序性知识图谱。",
        ],
    )
    add_body(
        doc,
        "在表示层面，Exposé 设想以 PKO 为目标本体：常见二元关系直接保存为边，复杂 n-ary procedural facts 则可通过 Execution 等中间节点 reification。该方法的目标是降低长文本的语义复杂度、提高关系一致性并减少幻觉 [S7][S16]。",
    )

    doc.add_heading("10.2.2 当前代码实际采用的方法", level=3)
    add_bullets(
        doc,
        [
            "从 video-to-text 之后开始：输入是工业 scene description、annotation-derived text 或 transcript-like text，不直接训练视频模型。",
            "执行确定性数据清洗，并生成 Raw 与 evidence-grounded Unified 两种输入；统一结构只重组原文支持事实。",
            "使用受 PKO 启发但经过简化的 Pydantic schema，配合 strict ontology/schema prompt、Instructor 和本地 Ollama 输出可校验对象。",
            "保留 one-shot baseline，同时实现 entity-first/relation-second layered extraction；第二阶段只能引用第一阶段已经抽取的 endpoint。",
            "增加 RelationCandidateGenerator、确定性规则和可选 LLM binary judge；LLM 只判断候选边是否有证据，不允许自由创造端点。",
            "把单个 scene 的抽取结果转换为 Property Graph，再执行节点规范化、方向校验、重复边合并、relation_origin 追踪与 source grounding。",
            "同时报告 Node/Edge P/R/F1、schema conformance、hallucination、raw/normalized 指标和 repeated-run stability。",
        ],
    )

    doc.add_heading("10.2.3 原计划与当前实现的逐项差异", level=3)
    add_table(
        doc,
        ["环节", "Exposé 原计划", "当前代码与状态", "调整原因"],
        [
            ["输入范围", "较长 instructional video transcripts 的预处理与切分", "处理 video-derived/annotation-derived industrial text；已完成", "当前可用数据已是文本和 scene annotation；先隔离视频识别误差，控制项目范围"],
            ["本体表示", "面向完整 PKO，并考虑 n-ary fact reification", "PKO-inspired 简化 Pydantic schema；部分实现", "小模型更容易稳定输出；Pydantic 可直接校验，Gold 也更容易一致标注"],
            ["WP3 依赖图", "显式构造 Action、时间、条件、共享资源的 procedural dependency graph", "已有 action sequence、temporal metadata 和 relation candidates，但无独立 dependency-graph data model；部分实现", "现有数据多为短 scene，条件/资源依赖标注不完整；先验证可评估关系"],
            ["WP4 分解", "对依赖图进行 tree/local-subtask decomposition", "entity-first/relation-second 两阶段抽取；部分实现", "关系端点依赖实体召回；两阶段更适合本地 7B/8B 模型，也便于单独定位实体和关系错误"],
            ["局部抽取", "每个子任务独立运行 ontology-guided prompt，形成 local subgraph", "对每个 scene 运行 strict schema one-shot 或 layered extraction；已完成简化版本", "当前 benchmark 是 scene-level，不具备稳定的长流程子任务边界；先建立可复现 baseline"],
            ["WP5 合并", "跨多个 local subgraph 做实体统一、属性冲突处理和全局合并", "单 scene Property Graph 的 stable ID、node merge、edge deduplication 和 direction correction；部分实现", "尚未生成真正独立的 local subgraphs，因此只能先实现图内规范化和去重"],
            ["关系生成", "LLM 在局部 prompt 中直接生成 triples", "candidate generation + deterministic validation + optional binary judge；新增改进", "端到端最佳 Edge F1 仅 0.1421；限制候选空间可减少自由造边并保留 evidence 审计轨迹"],
            ["评价", "以 Precision、Recall、F1 与图一致性比较 monolithic/decomposition", "增加 Node/Edge 分层指标、grounding/hallucination、raw/normalized 和 stability；已完成主要基础设施", "研究问题需要分别测量 correctness、completeness、hallucination 与 stability，且必须区分模型输出和后处理效果"],
        ],
        [1120, 2450, 2920, 2870],
        font_size=7.0,
        first_column_bold=True,
    )

    doc.add_heading("10.2.4 为什么采用当前的阶段性实现", level=3)
    add_labeled(doc, "数据与任务粒度", "Exposé 面向长转录文本，但当前正式工业 Gold 以 scene-level annotation 为主。直接实现 tree decomposition 会缺少可靠的 dependency labels 和局部子任务边界。")
    add_labeled(doc, "实验结果驱动", "关系抽取而不是 JSON 格式已成为主要瓶颈：工业端到端最佳 Edge F1 为 0.1421，而 32-scene Gold-entity candidate diagnostic F1 为 0.8620。差距表明实体遗漏、endpoint alignment 与 candidate coverage 必须先解决。")
    add_labeled(doc, "本地小模型约束", "7B/8B 模型在一次性生成完整图时出现 timeout、漏抽和格式漂移。先抽实体、再判断有限候选关系，可以降低单次输出复杂度并使错误可定位。")
    add_labeled(doc, "可审计与防幻觉", "当前 Gold 只允许 source-supported facts。Unified evidence、candidate_id、relation_origin 和 grounding report 能明确说明一条边来自 LLM、规则还是后处理，比自由生成局部子图更容易审核。")
    add_labeled(doc, "分阶段交付", "先证明单 scene 的清洗、schema extraction、graph construction 和 evaluation 可复现，再扩展 dependency graph、真实 local-subtask decomposition 和跨子图 conflict resolution，可减少多个未验证模块同时失败的风险。")
    add_callout(
        doc,
        "论文表述边界",
        "当前方法应写为 simplified layered extraction and evidence-constrained relation validation。完整 procedural dependency graph、tree decomposition 与 independent local-subgraph merge 仍属于下一阶段；不能把现有两阶段抽取和节点去重描述成这些方法已经完整实现。",
        risk=True,
    )
    add_source_note(
        doc,
        "docs/expose/expose_en_integrated.tex；tmp/proposal_review_20260731/proposal.txt；backend/pipeline/algorithm_experiment.py；backend/pipeline/relation_candidate_pipeline.py；backend/graph/property_graph.py；方法背景见 [S5][S6][S7][S16]。",
    )

    doc.add_page_break()
    doc.add_heading("10.3 数据清洗模块：确定性来源解析", level=2)
    add_code_figure(
        doc,
        code_figures["cleaning"],
        "代码图 1  工业来源记录的确定性清洗与报告生成",
        "Source excerpt for deterministic industrial data cleaning",
    )
    add_module_details(
        doc,
        role="把 IndEgo 等工业来源中的多个文本层统一为带来源追踪的清洁记录，同时报告空值、重复值和不完整字段。该层位于 LLM 之前。",
        input_output="输入为来源目录或记录列表；输出为 CleanFactoryRecord 列表和 FactoryCleaningReport。记录保留 source_path、layer、record_id 和 cleaning issues。",
        implementation="clean_factory_data_source 按文件类型调用确定性解析器，清除不可见字符和纯格式噪声，保留原始语义顺序；clean_indego_text_layers 对不同 annotation layer 使用同一规范。流程不调用 LLM，因此不会自动补充文本中没有的动作、工具或因果关系。",
        experiment="医院与工业实验表明，清洗模块主要减少格式错误和无效字段。它提高输入可用性，但本身不会提高关系语义 recall；清洗后的事实仍需 schema 抽取和 Gold 比较。",
        references="数据字段设计来自 IndEgo annotation/metadata [S14]；清洗和报告逻辑为本项目实现。",
        boundary="当前清洗面向项目已知 JSON/CSV/text layer。对新的工业来源仍需新增 adapter；不把人工推理或 AI 补写内容写入 cleaned record。",
    )
    add_source_note(doc, "backend/cleaning/factory_data.py；backend/datasets/indego_adapter.py")

    doc.add_heading("10.4 统一文本模块：只重组原文支持事实", level=2)
    add_code_figure(
        doc,
        code_figures["unified"],
        "代码图 2  UnifiedTextRecord 构建与 evidence 校验",
        "Source excerpt for grounded unified-text construction",
    )
    add_module_details(
        doc,
        role="将同一段 raw text 重组为 scene、action sequence、tools/objects、parameters、quality result 和 uncertainty 等稳定字段，供小模型读取。",
        input_output="输入为 raw_text 与 GroundedEntry 列表；输出为 UnifiedTextRecord、JSON 表示和可直接放入 prompt 的统一文本。",
        implementation="_validate_evidence 对每个结构化条目执行 normalized substring 检查：evidence 必须存在于 raw_text。build_unified_text 只组装已验证字段，并保留 segment ID 和 evidence；不通过校验的新增事实会被拒绝。",
        experiment="早期 5 场景实验中，Raw schema success 为 3/5，Unified 为 5/5。该效果说明统一结构改善了输出可解析性；但 Unified Relation F1 曾为 0，因此报告不把格式稳定性等同于关系正确性。",
        references="输入形式作为实验变量的设计由本项目提出；ontology-guided text-to-KG 的评价背景参考 Text2KGBench [S5] 与 KaLLM maintenance KG [S6]。",
        boundary="MVP 中统一结构由人工或规则辅助生成。自动 LLM preprocessing 若未通过 evidence 校验，可能引入幻觉，不进入 Gold 或正式 Unified input。",
    )
    add_source_note(doc, "backend/preprocessing/unified_text.py；docs/results/local_experiment_2026-05-26_r1.md")

    doc.add_page_break()
    doc.add_heading("10.5 Schema 与本地结构化 LLM 抽取", level=2)
    add_code_figure(
        doc,
        code_figures["schema"],
        "代码图 3  EgocentricVideoExtraction 的 Pydantic 输出合同",
        "Pydantic extraction contract for entities and procedural relations",
    )
    add_module_details(
        doc,
        role="定义 LLM 可以生成的实体和关系容器，是整个 pipeline 的可执行 schema。核心实体包括 Scene、Action、Tool、SceneObject、Procedure 与 ProcessParameter。",
        input_output="输入是模型返回的 JSON-like payload；输出是通过 Pydantic 校验的 EgocentricVideoExtraction，内含实体列表、关系列表和 provenance。",
        implementation="每类实体和关系由强类型 Pydantic model 表示；关系端点通过稳定 ID 连接。Pydantic 在对象构造时验证字段类型、必填项和枚举约束，从而把不符合 schema 的输出变成显式 error。",
        experiment="Schema validation 是所有真实 LLM 实验的第一道门。Unified input 和 JSON repair 降低了 schema failure；但通过 schema 只说明结构合法，并不说明事实与原文一致。",
        references="Pydantic model/validation API [S1]；procedural categories 受 PKO [S7] 启发，但项目采用便于 LLM 抽取和评估的简化 schema。",
        boundary="当前 schema 没有完整复刻 PKO，也没有覆盖所有工业过程关系。新增关系时必须同步更新 ontology domain/range、property graph、evaluation 和 Gold。",
    )
    add_source_note(doc, "backend/schemas/egocentric_video.py；backend/schemas/ontology.py")

    add_code_figure(
        doc,
        code_figures["llm"],
        "代码图 4  Ollama、OpenAI-compatible client 与 Instructor 的连接方式",
        "Local structured extraction client using Ollama and Instructor",
    )
    add_module_details(
        doc,
        role="统一封装本地模型调用，把 prompt、输入文本和 Pydantic response model 送入 Ollama，并返回已经过结构校验的 Python 对象。",
        input_output="输入为文本、response_model、模型名、prompt、temperature 和 timeout；输出为指定 Pydantic 类型，失败时抛出可记录的异常。",
        implementation="OpenAI Python client 指向 Ollama 的 /v1 兼容端点，Instructor 使用 response_model 执行结构化解析和有限重试。temperature 默认设为 0，以减少随机变化。",
        experiment="真实 LLM 长实验能够统一切换 llama3.1:8b、qwen2.5:7b 等模型，并把 timeout/parse failure 记录为 error。模型对结果有影响，但当前样本不足以形成最终模型排名。",
        references="OpenAI Python client [S3]、Ollama OpenAI compatibility [S4]、Instructor structured outputs [S2]。这里属于公开 API 的直接使用。",
        boundary="Instructor 重试只能修复部分结构错误，不能证明事实正确。Ollama 服务、模型版本、上下文窗口和硬件状态都会影响运行时间与可重复性。",
    )
    add_source_note(doc, "backend/llm/client.py")

    add_code_figure(
        doc,
        code_figures["schema_prompt"],
        "代码图 5  Extractor 中的 ontology/schema prompt 约束",
        "Schema guidance and exact relation-shape constraints in the extractor",
    )
    add_module_details(
        doc,
        role="在模型调用前把允许的实体、关系名称、端点方向和禁止事项写入 prompt，减少 relation label 漂移和无证据事实。",
        input_output="输入为 response model 与待抽取文本；输出为由系统 prompt 和 schema guidance 约束的结构化调用。",
        implementation="_schema_guidance 根据目标 Pydantic model 生成实体字段、exact relation shapes 和 domain/range 提示，明确要求 only extract explicitly supported facts。",
        experiment="医院 stress 实验显示 strict schema prompt 优于宽松 prompt；因此保留为正式 pipeline 的基础约束。关系 F1 仍低，说明 prompt 约束无法单独解决端点漏召回。",
        references="Text2KGBench 的 ontology-driven generation [S5]、van Cauter 与 Yakovets 的 ontology-guided triplet extraction [S6]；实现代码为本项目适配。",
        boundary="Prompt 约束属于软约束。最终仍需 Pydantic validation、domain/range validation 和 source grounding，不能只相信模型遵守指令。",
    )
    add_source_note(doc, "backend/extraction/extractor.py")

    doc.add_heading("10.6 实验编排与逐层抽取", level=2)
    add_code_figure(
        doc,
        code_figures["runner"],
        "代码图 6  一次实验运行的抽取、图构建与验证顺序",
        "Experiment runner sequence for extraction graph construction and validation",
    )
    add_module_details(
        doc,
        role="保证 raw/unified、one-shot/layered、不同模型和重复运行使用同一套保存与评价流程，避免实验条件之间出现隐性差异。",
        input_output="输入为 benchmark item 和 ExperimentCondition；输出为 RunRecord/checkpoint，包括 extraction、graph、validation issues、metrics、runtime 和 error。",
        implementation="_run_once 先选定条件文本，再调用 Extractor，随后转换为 PropertyGraph；hallucination validation 始终对照 raw source text，而不是对照模型整理后的文本。每次运行立即落盘，长实验中断后仍能审计。",
        experiment="该层支撑 EASG 10-scene、工业 5-scene 和 repeated-run 实验。它不会改善模型本身，但保证错误、超时和失败运行不会被静默丢弃。",
        references="实验记录结构由本项目实现；评价维度与 Text2KGBench [S5]、KaLLM [S6] 的 fact extraction/ontology conformance 思路一致。",
        boundary="不同历史 runner 的字段曾不完全一致，汇总时需按 experiment mode 区分。正式复现实验还应固定 prompt hash、模型 digest 和数据 split。",
    )
    add_source_note(doc, "backend/pipeline/experiment_runner.py；backend/evaluation/experiment.py")

    add_code_figure(
        doc,
        code_figures["layered"],
        "代码图 7  Entity-first / relation-second 两阶段抽取",
        "Two-stage layered extraction implementation",
    )
    add_module_details(
        doc,
        role="把一次性生成整张图拆成实体抽取和关系抽取两个阶段，降低小模型单次任务复杂度，并使关系错误可以单独诊断。",
        input_output="第一阶段输入文本并输出 entity inventory；第二阶段输入原文和实体列表，输出 relation result；最后合成为 EgocentricVideoExtraction。",
        implementation="_two_stage_extract 使用独立 Pydantic response models；关系 prompt 只能引用已抽取 endpoint。minimal 版本进一步减少字段，使候选关系实验可以控制生成空间。",
        experiment="Layered 条件通常更容易运行，但工业实验中最佳 Edge F1 来自 raw+one-shot 0.1421，unified+layered 的优势主要体现在 Node F1 0.6019。由此不能声称 layered 已稳定提高关系质量。",
        references="逐层分解受复杂任务 decomposition 思想启发；当前代码只实现 entity-first/relation-second，不是完整 dependency graph/tree decomposition。",
        boundary="如果第一阶段漏掉实体，第二阶段无法生成相关边；entity recall 构成 relation recall 的上限之一。",
    )
    add_source_note(doc, "backend/pipeline/algorithm_experiment.py；docs/results/industrial_full_after_relation_fix_2026-07-05.md")

    doc.add_page_break()
    doc.add_heading("10.7 关系候选生成与二元判断", level=2)
    add_code_figure(
        doc,
        code_figures["candidate"],
        "代码图 8  RelationCandidateGenerator 的候选生成入口",
        "Relation candidate generation interface and deterministic context assembly",
    )
    add_module_details(
        doc,
        role="先根据已知实体、文本共现、动作顺序和 annotation metadata 生成有限候选边，再交给规则或评分器判断，减少 LLM 自由造边。",
        input_output="输入为文本、实体、scene/action sequence 和可选 temporal metadata；输出为带 candidate_id、source_id、relation、target_id、evidence/context 的 RelationCandidate 列表。",
        implementation="Generator 针对 USES_TOOL、ACTS_ON、BEFORE、PART_OF、OBSERVED_IN、WARNING_FOR 使用不同规则；端点必须来自已知实体。候选去重后保留 generation_reason，便于计算 coverage recall。",
        experiment="32 场景 Gold-entity 诊断总体 F1 0.8620，但 USES_TOOL Recall 仅 0.2979。该结果说明确定性候选对 BEFORE 很强，对共享工具和跨句工具证据覆盖不足。",
        references="受 LangChain LLMGraphTransformer 的 allowed nodes/relationships 与 strict filtering [S10]、GLiREL 候选关系分类接口 [S11] 启发；具体生成规则为本项目实现。",
        boundary="Gold-entity diagnostic 直接使用人工实体，不是端到端 LLM 分数。Generator 未生成的边，后续 scorer 无法恢复。",
    )
    add_source_note(doc, "backend/pipeline/relation_candidate_pipeline.py；docs/results/industrial_reviewed_gold_combined_relation_diagnostic_2026-08-01.md")

    add_code_figure(
        doc,
        code_figures["binary_judge"],
        "代码图 9  LLM binary judge 只判断候选边是否受支持",
        "Binary relation judge restricted to existing candidate IDs",
    )
    add_module_details(
        doc,
        role="让 LLM 对每条已存在候选边回答 supported/not supported，而不是开放式生成 relation tuple。",
        input_output="输入为 candidate_id、候选端点、relation type 和 source context；输出为 is_supported、confidence、evidence_text 与 reason。",
        implementation="Prompt 强制返回候选 ID，并在 scorer 侧检查 evidence 是否来自输入。判断器不能修改端点或 relation label；被拒绝的候选进入 validation report。",
        experiment="5-scene 消融中 Edge F1 从 deterministic 0.0476 提升到 0.0564，但平均耗时增至 81.84 秒。Judge 主要提高 precision，未解决候选漏召回。",
        references="GLiREL 的 relation classification 设计 [S11] 与 LangChain strict relationship filtering [S10]；当前实验使用本地 LLM binary judge，未实际运行 GLiREL 模型。",
        boundary="GLiRELScorer 目前只是可选接口槽位，不能在论文中写成已完成对照实验。Binary judge 的置信度也不是经过校准的概率。",
    )
    add_source_note(doc, "backend/pipeline/relation_candidate_pipeline.py；docs/results/relation_candidate_ablation_2026-07-07.md")

    doc.add_page_break()
    doc.add_heading("10.8 Property Graph 构建", level=2)
    add_code_figure(
        doc,
        code_figures["graph"],
        "代码图 10  节点合并、稳定 ID 与边去重",
        "In-memory property graph node merge and edge deduplication",
    )
    add_module_details(
        doc,
        role="把 Pydantic extraction 转为统一 GraphNode/GraphEdge，合并重复节点，保留关系属性和 provenance，并支持基础查询与 Cypher 导出。",
        input_output="输入为 EgocentricVideoExtraction；输出为内存 PropertyGraph，其中 nodes/edges 使用稳定 ID 和标准 label。",
        implementation="节点先按 stable ID 查找，再使用 normalized name/token overlap 处理可确定别名；边以 source、relation、target 作为去重键。merge 时合并 properties 和 provenance，不覆盖已有证据。",
        experiment="该层使不同抽取条件可以转换为相同图表示并计算 Node/Edge F1。节点合并减少表面重复，但语义错误分类仍需 post-processing 或人工 Gold 判断。",
        references="设计参考 Neo4j LLM Graph Builder [S8] 与 LlamaIndex Property Graph 示例 [S9]；当前实现是本项目的轻量内存 MVP，不依赖 Neo4j/LlamaIndex 运行时。",
        boundary="尚未接入持久数据库，也没有完整图事务和大规模索引。token overlap 合并必须保守，否则可能把不同工业零件错误合并。",
    )
    add_source_note(doc, "backend/graph/property_graph.py")

    doc.add_heading("10.9 工业后处理与关系来源追踪", level=2)
    add_code_figure(
        doc,
        code_figures["postprocess"],
        "代码图 11  节点规范化、关系别名、方向修正与重复边合并",
        "Industrial KG post-processing with relation origin tracking",
    )
    add_module_details(
        doc,
        role="在模型输出后修复可确定的格式、类别和方向错误，同时记录每条边来自 LLM、evidence fallback、sequence fallback 或 post-processing。",
        input_output="输入为 IndustrialNode/IndustrialEdge、source support index 和 schema；输出为 normalized graph、统计信息、merge log 和 validation issues。",
        implementation="节点 label/name 先规范化；工具分类使用可维护词表和 use/with/by using 等上下文；relation alias 映射到标准名；若 schema 能唯一判断反向边则翻转；最后按三元组去重并检查 evidence support。",
        experiment="医院 alias/reverse stress 可把 normalized F1 从 0 修复到 1.0，证明格式后处理有效；但 missing fact 条件无法恢复。工业产品中 fallback 提高可用边数，raw 与 normalized 指标仍分开报告。",
        references="Ontology domain/range validation 参考 Text2KGBench [S5] 与 KaLLM [S6]；具体词表、relation_origin 和安全修正规则为本项目实现。",
        boundary="后处理不能伪装成模型原始能力。只有能由 schema 和原文唯一确认的边才自动修正；无法确认的内容标为 unsupported。",
    )
    add_source_note(doc, "backend/pipeline/industrial_text_to_kg.py；config/industrial_tool_terms.txt")

    doc.add_page_break()
    doc.add_heading("10.10 评价：P/R/F1、Hallucination 与 Stability", level=2)
    add_code_figure(
        doc,
        code_figures["metrics"],
        "代码图 12  Precision、Recall 与 F1 的统一计算",
        "Shared precision recall and F1 implementation",
    )
    add_module_details(
        doc,
        role="为节点、属性和关系提供统一的 TP/FP/FN 指标对象，避免不同实验脚本各自计算导致口径漂移。",
        input_output="输入为 true_positives、false_positives、false_negatives；输出为 PRF1，包含 precision、recall、f1 和计数。",
        implementation="Precision=TP/(TP+FP)，Recall=TP/(TP+FN)，F1 为调和平均；空分母显式返回 0。evaluators 先把 Gold 与 prediction 规范化为可比较 fact sets，再调用 compute_prf1。",
        experiment="工业端到端最佳 Node F1 为 0.6019，最佳 Edge F1 为 0.1421；32-scene 0.8620 是 Gold-entity 关系诊断。两者输入条件不同，必须分别陈述。",
        references="指标定义采用标准 information extraction 评价；Text2KGBench [S5] 与 KaLLM [S6] 使用相同的 Precision/Recall/F1 评价背景。",
        boundary="F1 不评价证据是否可信，也不反映重复运行波动，因此必须与 grounding/hallucination 和 stability 指标联合使用。",
    )
    add_source_note(doc, "backend/evaluation/metrics.py；backend/evaluation/evaluators.py")

    add_code_figure(
        doc,
        code_figures["hallucination"],
        "代码图 13  Ontology conformance 与文本 grounding 检查",
        "Ontology and grounding validation for hallucination analysis",
    )
    add_module_details(
        doc,
        role="检测非法关系、错误 domain/range、未知端点，以及 subject/object/evidence 缺乏原文支持的情况。",
        input_output="输入为 extraction、ontology spec 和 source_text；输出为 validation issues、ontology conformance rate 与 hallucination rate 分项。",
        implementation="每条边先检查 relation 是否存在，再验证 source/target label 是否符合 schema；节点名和 evidence 通过 normalized grounding 与 source text 比较。每个 issue 保留类别、fact 和说明。",
        experiment="该模块帮助区分格式合法与事实有依据。项目已能输出 unsupported nodes/edges；但 32 场景正式 hallucination 汇总尚未全部完成，因此不能给出最终总体幻觉率。",
        references="Hallucination/ontology conformance 评价参考 Text2KGBench [S5] 和 maintenance short-text KG 工作 [S6]；validation 实现为本项目代码。",
        boundary="字符串 grounding 不能识别所有同义表达，也不能自动判断复杂隐含关系。正式 Gold 仍以人工确认的 source-supported facts 为准。",
    )
    add_source_note(doc, "backend/evaluation/ontology_validation.py")

    add_code_figure(
        doc,
        code_figures["stability"],
        "代码图 14  重复运行图重合度与波动计算",
        "Repeated-run graph overlap and variation metrics",
    )
    add_module_details(
        doc,
        role="衡量同一输入在重复抽取时生成节点集合和关系集合的一致程度，用于 operationalize stability。",
        input_output="输入为同一 condition 的多个 RunRecord；输出为 pairwise node/relation Jaccard、graph overlap、平均值和 variation。",
        implementation="先把每次图转换为 normalized node/edge fact sets，再计算两两 Jaccard；graph overlap 综合节点与边，variation=1-overlap。错误运行单独计数。",
        experiment="部分实验中 overlap 可达到 1.0，但 Edge F1 仍为 0，说明模型可能稳定地生成同一种错误。Stability 不能代替 correctness 和 completeness。",
        references="稳定性的操作化定义由本项目实验设计确定；与 repeated extraction agreement 的通用思想一致。",
        boundary="当前完整 32-scene、多模型、多次端到端 stability test 尚未完成；重复次数过少时统计结论不稳。",
    )
    add_source_note(doc, "backend/evaluation/stability.py；backend/evaluation/reporting.py")

    doc.add_page_break()
    doc.add_heading("10.11 Gold 与公开数据 Adapter", level=2)
    add_code_figure(
        doc,
        code_figures["gold"],
        "代码图 15  人工审核 Gold 的端点、签名与重复关系校验",
        "Validation of reviewed industrial Gold entities and relations",
    )
    add_module_details(
        doc,
        role="把人工审核后的工业 scene 转为统一 benchmark item，并在加载时阻止悬空端点、重复 ID、非法 relation signature 和无 source evidence 的条目进入正式实验。",
        input_output="输入为 reviewed JSONL；输出为 BenchmarkDataset/EgocentricVideoExtraction 与 validation report。",
        implementation="Loader 建立 node ID 索引，逐条检查 relation source/target 是否存在和 label 是否符合 schema；按 video_id 保存 split metadata，但 reviewer notes 不进入模型 Unified input。",
        experiment="V2+V3 合并后形成 32 scenes、30 video IDs、674 entities、828 relations，用于关系候选诊断。正式 train/test 划分必须按 video_id，避免同一视频片段泄漏。",
        references="EASG graph annotation 格式与人工图来源 [S12][S13]、IndEgo 工业 annotation/metadata [S14]；Reviewed Gold 的人工修订与 loader 为项目内部工作。",
        boundary="32 scenes 仍属于开发/回归规模，不足以支撑强泛化结论。候选自动生成的数据只能叫 silver_candidate，必须人工审核后才是 Gold。",
    )
    add_source_note(doc, "backend/datasets/industrial_reviewed_gold.py；backend/datasets/easg_adapter.py；backend/datasets/indego_adapter.py")

    doc.add_heading("10.12 一键产品入口与输出格式", level=2)
    add_code_figure(
        doc,
        code_figures["product"],
        "代码图 16  批量读取 .txt 并调用完整工业 Text-to-KG pipeline",
        "Product-facing batch text-to-KG entry point",
    )
    add_module_details(
        doc,
        role="提供组员可以直接运行的成品入口：把输入目录中的工业 .txt 文件逐个转换为知识图谱，并为每个文件建立独立结果目录。",
        input_output="输入为一个或多个 UTF-8 .txt 文件以及模型/timeout 参数；输出为每个文本的 kg_result.json、nodes.csv、edges.csv、graph.cypher 和 summary.md。",
        implementation="Runner 对文件排序，调用 build_industrial_kg_from_text，捕获单文件异常并继续批处理；所有结果记录输入文件、模型、节点边统计和 validation summary。",
        experiment="随机 pump assembly 文本和组员电脑复现均能生成节点与边。14/14 release tests 通过；旧医院路径失败与工业成品链路分离。",
        references="命令行与批处理实现为本项目代码；导出结构参考 Property Graph/Neo4j 的常见节点边表示 [S8][S9]。",
        boundary="Product V1 的关系包含 llm_extracted 与 evidence/sequence fallback，必须查看 relation_origin。输出可导入图库，但当前包不自动启动 Neo4j。",
    )
    add_source_note(doc, "scripts/run_batch_industrial_texts_to_kg.py；scripts/run_custom_industrial_text_to_kg.py")

    add_code_figure(
        doc,
        code_figures["export"],
        "代码图 17  JSON、CSV、Cypher、摘要和批次索引导出",
        "Export of graph artifacts and batch audit records",
    )
    add_module_details(
        doc,
        role="把同一图以机器可读、表格查看、数据库导入和人工审计四种形式保存，使结果可复现和可交换。",
        input_output="输入为 pipeline result 和 output directory；输出为完整图 JSON、节点/边 CSV、Cypher、单文件 summary 及 RUN_SUMMARY。",
        implementation="JSON 保留全部 properties、provenance、validation issues 和 relation_origin；CSV 展平关键字段；Cypher 生成节点与关系语句；batch index 汇总成功、失败与路径。",
        experiment="该层让组员无需打开终端分析内部对象，也可以直接检查图和报告。它不改变模型质量，只提高交付、复现和错误追踪能力。",
        references="Property Graph 导出设计参考 Neo4j [S8]；文件组织、审计字段和批次报告为本项目实现。",
        boundary="CSV 会丢失部分嵌套结构，研究审计应以 kg_result.json 和原始 run record 为准。Cypher 在导入生产数据库前仍需检查转义和目标版本。",
    )
    add_source_note(doc, "scripts/run_batch_industrial_texts_to_kg.py")

    doc.add_heading("10.13 代码模块与实验结论对照", level=2)
    add_table(
        doc,
        ["模块/措施", "验证它的实验", "观察到的效果", "决策"],
        [
            ["Unified text", "E1/E9/E10", "提高 schema success/节点可用性；Edge 不稳定", "保留，不能单独声称提升 KG 质量"],
            ["Strict schema prompt", "医院 prompt stress", "严格约束优于宽松输出", "保留为默认"],
            ["Layered extraction", "P2/E9/E10", "更可运行；关系未稳定超过 one-shot", "保留为实验条件，继续优化"],
            ["JSON repair", "EASG 10-scene", "解析错误降到 0", "保留安全修复，失败仍记 error"],
            ["Post-processing", "医院 alias/reverse stress", "可修复格式、别名、方向；不能补漏事实", "保留并分报 raw/normalized"],
            ["Candidate generator", "E11/E14", "BEFORE 强，USES_TOOL recall 弱", "保留，优先扩展候选 recall"],
            ["LLM binary judge", "E12", "Edge F1 小幅增加，成本明显增加", "可选，不作为默认全量步骤"],
            ["GLiREL scorer", "尚无实验", "只有接口，无真实对照结果", "Not started"],
            ["Hallucination validation", "后处理/回归", "可报告 unsupported facts", "保留，扩大正式统计"],
            ["Stability metrics", "repeated-run pilots", "可发现输出波动；稳定不等于正确", "与 F1 联合报告"],
        ],
        [2000, 2200, 3100, 2060],
        font_size=7.4,
        first_column_bold=True,
    )

    doc.add_page_break()
    doc.add_heading("11. 遇到的问题与对应改进", level=1)
    add_table(
        doc,
        ["问题", "实验表现", "改进措施", "当前效果/剩余风险"],
        [
            ["研究范围过大", "容易被理解为直接视频理解", "范围固定在 video-derived text 之后", "原始视频模型仍不在 MVP"],
            ["简单任务无区分度", "P3/P4 初始条件均满分", "增加 noisy、dense、loose prompt stress", "参数敏感，需固定实验协议"],
            ["One-shot 任务过重", "医院全图 180s timeout", "entity-first、relation-second layered extraction", "更可运行，但关系 recall 仍低"],
            ["JSON 不稳定", "EASG raw+layered 多次解析失败", "安全 JSON repair，失败仍记 error", "10-scene 达到 0 error"],
            ["关系 label 解析 bug", "ACTS_ON 被清洗为 ACTS ON", "grouped payload + normalize_relation_type", "Edge F1 提升，但仍不高"],
            ["节点 ID/别名噪声", "raw F1 可为 0", "canonical ID、alias index、node normalization", "格式错误可恢复，漏抽不可恢复"],
            ["反向与重复边", "模型方向错误、重复输出", "schema domain/range 纠正和 merge", "必须标为 postprocessed"],
            ["幻觉检查不足", "unsupported IDs/edges 混入图", "source grounding + ontology validation", "32 场景正式 hallucination 汇总未完成"],
            ["Stability 被误读", "overlap 1.0 但 Edge F1 0", "与 P/R/F1、errors 联合报告", "完整重复端到端 WP8 仍待完成"],
            ["候选边漏召回", "USES_TOOL recall 0.2979", "共享工具、跨句证据、endpoint alignment", "当前首要技术任务"],
            ["LLM judge 成本高", "Edge F1 仅 +0.0088，81.84s", "先缩短候选列表，再做 binary judge", "GLiREL 对照尚未执行"],
            ["数据域差异", "EASG 优于工业 benchmark", "医院/EASG/工业角色分离", "最终论文以工业 Gold 为主"],
            ["Gold 规模不足", "早期只有 5 或 12 scenes", "人工 V3 合并到 32 scenes", "仍建议抽样二审和更多独立视频"],
            ["自动报告措辞错误", "工业 smoke 写成 EASG Gold", "报告模板按 dataset_name 生成", "旧结果保留并在本报告标注"],
        ],
        [1850, 2350, 2750, 2410],
        font_size=7.6,
        first_column_bold=True,
    )

    doc.add_page_break()
    doc.add_heading("11.1 改进措施的方法来源", level=2)
    add_table(
        doc,
        ["措施", "方法原理", "来源/项目实现"],
        [
            ["Ontology-guided extraction", "给定允许实体、关系及 domain/range，减少自由关系命名", "Text2KGBench；van Cauter & Yakovets 2024；项目 schema"],
            ["Layered extraction", "把实体识别与关系判断拆开，降低单次输出复杂度", "项目 pipeline；受任务分解思想启发，但不是完整 tree decomposition"],
            ["Candidate relation validation", "只在有限候选上判断支持性，禁止模型创造端点", "RelationCandidateGenerator/Scorer/Report"],
            ["Grounding validation", "每个事实必须能回到 source evidence", "项目 ontology_validation 与 reviewed Gold 规则"],
            ["Raw/normalized dual metrics", "分离模型能力与确定性后处理效果", "医院 post-processing long benchmark"],
            ["Repeated-run stability", "用集合 Jaccard 与指标方差量化重复一致性", "backend/evaluation/stability.py；P6/EASG experiments"],
        ],
        [2500, 3550, 3310],
        font_size=8.0,
        first_column_bold=True,
    )

    doc.add_page_break()
    doc.add_heading("12. 模块保留决策", level=1)
    add_table(
        doc,
        ["模块", "决策", "实验依据", "论文表述边界"],
        [
            ["Strict schema prompt", "保留", "Loose 条件关系质量明显下降", "ontology constraint"],
            ["Unified text structure", "保留", "提高 schema/ID 可用性；工业 Node 结果有提升", "不声称总能提高 F1"],
            ["Layered extraction", "保留", "one-shot 曾超时；layered 更可运行", "简化任务拆分，不是完整 decomposition"],
            ["JSON repair", "保留", "EASG 10-scene 0 error", "只修语法，不补事实"],
            ["Node normalization", "保留", "alias_noise normalized 1.0", "报告 raw 与 normalized"],
            ["Relation direction/merge", "保留", "反向/重复噪声可修复", "标注 relation_origin"],
            ["Hallucination validation", "保留", "unsupported IDs/edges 可检测", "需要 32-scene 正式统计"],
            ["Candidate generator interface", "保留并继续开发", "架构可用，但 deterministic F1 未提升", "重点提高 recall"],
            ["LLM binary judge", "条件保留", "Edge F1 0.0476→0.0564", "耗时高，只是 precision filter"],
            ["One-shot full graph", "Baseline only", "关系有时较高，但 timeout/格式风险", "不作为最终主流程"],
            ["EASG simulation", "实验框架验证", "125000 模拟 runs", "不得当作真实模型结果"],
        ],
        [2200, 1400, 3260, 2500],
        font_size=8.0,
        first_column_bold=True,
    )

    doc.add_heading("12.1 当前真正完成的系统", level=2)
    add_numbered(
        doc,
        [
            "Data Cleaning Module：raw/unified 输入、evidence 保留、可保存 JSON 和 LLM 文本。",
            "Text-to-KG Module：one-shot/layered 抽取、候选关系、属性图、节点合并、边来源、JSON/CSV/Cypher 导出。",
            "Evaluation Layer：Node/Edge P/R/F1、ontology/grounding/hallucination 检查、稳定性、raw/normalized 指标。",
            "Reviewed Gold Adapter：V2、V3 及 combined 32-scene 数据可重复加载和诊断。",
        ],
    )
    add_callout(
        doc,
        "尚未完成",
        "完整 procedural dependency graph、基于 tree decomposition 的局部子任务、独立局部子图的冲突合并、32 场景重复端到端 LLM 与正式 hallucination/stability 汇总仍未完成。",
        risk=True,
    )

    doc.add_page_break()
    doc.add_heading("13. 最终研究成果与下一步", level=1)
    doc.add_heading("13.1 已经得到的可靠结论", level=2)
    reliable_conclusions = [
        "统一结构最稳定的作用是减少格式、ID 和 schema 不一致，而不是自动保证关系 F1 提高。",
        "Layered extraction 可以降低一次性任务复杂度，但关系策略本身仍需单独优化。",
        "工业端到端的最佳 Node F1 为 0.6019；最佳 Edge F1 为 0.1421，说明关系仍是系统瓶颈。",
        "在 Gold entities 正确时，32-scene candidate diagnostic F1 可达到 0.8620，表明端到端低分还受到实体遗漏与 endpoint mismatch 影响。",
        "BEFORE 已较成熟，ACTS_ON 基本可用，USES_TOOL 的 candidate recall 是当前最明显短板。",
        "Post-processing 对别名、方向和重复问题有效，但不能替代上游 completeness。",
        "稳定性必须与 correctness 分开；模型可以稳定地重复错误图。",
    ]
    for index, conclusion in enumerate(reliable_conclusions, 1):
        add_body(doc, f"{index}. {conclusion}", after=3)

    doc.add_heading("13.2 优先级计划", level=2)
    add_table(
        doc,
        ["优先级", "任务", "验收标准"],
        [
            ["P0", "审核 20-32 scenes 的 relation Gold 抽样二审", "证据、方向、端点和完整性一致"],
            ["P1", "改进 USES_TOOL candidate generator", "Recall 明显高于 0.2979，且 precision 不崩溃"],
            ["P2", "完善 ACTS_ON endpoint alias 和跨句证据", "降低 138 条 missed Gold 中的可修复部分"],
            ["P3", "缩短候选后重跑 LLM binary judge", "Edge F1 提升大于当前 0.0088，runtime 可控"],
            ["P4", "运行 32 scenes × 至少 2 repeats × 4 conditions", "完整 Node/Edge、errors、hallucination、stability 表"],
            ["P5", "接入处理好的视频 metadata", "时间段、narration、object labels 可追踪到 evidence"],
        ],
        [1200, 3900, 4260],
        font_size=8.2,
        first_column_bold=True,
    )

    doc.add_page_break()
    doc.add_heading("13.3 可用于论文的最终表述", level=2)
    add_callout(
        doc,
        "建议表述",
        "Experiments indicate that unified text representation and layered extraction improve output-format reliability and, under some industrial conditions, entity-level F1. However, relation quality remains limited by candidate coverage and endpoint alignment. Gold-entity diagnostics show that deterministic relation candidates can recover most BEFORE and ACTS_ON edges, while USES_TOOL remains the primary recall bottleneck.",
    )
    add_body(
        doc,
        "中文：实验表明，统一文本表示与逐层抽取可以改善输出格式可靠性，并在部分工业条件下提高实体层 F1。然而，关系质量仍受候选覆盖和端点对齐限制。使用人工 Gold entities 的诊断说明，当前规则已能覆盖大多数 BEFORE 和 ACTS_ON，但 USES_TOOL 仍是主要召回瓶颈。",
    )

    doc.add_page_break()
    doc.add_heading("附录 A：全部结果目录索引", level=1)
    add_body(doc, "以下索引自动读取当前 results/ 目录。它包含正式实验、smoke、回归、开发诊断和未完成运行。只有“主要证据”进入正文核心结论。")
    result_dirs = sorted(path.name for path in (ROOT / "results").iterdir() if path.is_dir())
    appendix_rows: list[list[str]] = []
    for index, name in enumerate(result_dirs, 1):
        category, use = classify_result_dir(name)
        appendix_rows.append([f"R{index:02d}", name, category, use])
    add_table(doc, ["编号", "结果目录", "类型", "报告用途"], appendix_rows, [700, 4550, 1800, 2310], font_size=7.2, first_column_bold=True)
    add_source_note(doc, f"results/；共发现 {len(result_dirs)} 个结果目录。")

    doc.add_page_break()
    doc.add_heading("附录 B：主要报告与数据文件", level=1)
    source_rows = [
        ["早期工业", "docs/results/local_experiment_2026-05-26_r1.md"],
        ["医院 Pilot", "docs/results/hospital_gold_pilot_2026-06-04.md"],
        ["医院 Stress", "docs/results/hospital_gold_pilot_p134_stress_zh.md"],
        ["医院 Long", "results/real_llm_long_2026-06-19/summary.md"],
        ["医院后处理", "results/hospital_text_processing_long_2026-06-19/summary.md"],
        ["EASG Full Simulation", "results/easg_module_experiment_full_2026-06-20/summary.md"],
        ["EASG 10-scene", "docs/results/easg_real_llm_lightweight_realtest_10scene_2026-06-20.md"],
        ["工业初始 Full", "docs/results/industrial_real_llm_lightweight_final_2026-06-20.md"],
        ["工业关系修复", "docs/results/industrial_full_after_relation_fix_2026-07-05.md"],
        ["工业复现", "docs/results/industrial_full_after_relation_fix_rerun2_2026-07-05.md"],
        ["Candidate Pipeline", "docs/results/industrial_candidate_pipeline_eval_2026-07-05.md"],
        ["Relation Ablation", "docs/results/relation_candidate_ablation_2026-07-07.md"],
        ["Reviewed V2 Diagnostic", "docs/results/industrial_reviewed_gold_v2_relation_diagnostic_2026-08-01.md"],
        ["Combined Diagnostic", "docs/results/industrial_reviewed_gold_combined_relation_diagnostic_2026-08-01.md"],
        ["Combined Gold", "data/industrial_gold_reviewed/industrial_gold_combined_v2_v3_reviewed.jsonl"],
    ]
    add_table(doc, ["主题", "文件"], source_rows, [2500, 6860], font_size=8.0, first_column_bold=True)

    doc.add_heading("附录 C：外部方法参考", level=1)
    references = [
        "Mihindukulasooriya et al. (2023). Text2KGBench: A Benchmark for Ontology-Driven Knowledge Graph Generation from Text.",
        "van Cauter and Yakovets (2024). Ontology-guided Knowledge Graph Construction from Maintenance Short Texts. KaLLM 2024.",
        "Carriero et al. (2025). Procedural Knowledge Ontology (PKO).",
        "Zhou et al. (2025). An Approach for Systematic Decomposition of Complex LLM Tasks. 该文只支持计划中的完整 decomposition，不代表项目已经完成 tree decomposition。",
    ]
    for index, reference in enumerate(references, 1):
        add_body(doc, f"[R{index}] {reference}")

    doc.add_page_break()
    doc.add_heading("附录 D：代码、方法与数据来源索引", level=1)
    add_body(
        doc,
        "以下来源编号与第 10 章一致。直接依赖表示仓库通过公开 API 调用该软件；方法参考表示项目吸收其设计思想并自行实现；数据来源表示 adapter 对接其 annotation/metadata。",
    )
    doc.add_heading("D.1 直接代码依赖与 API", level=2)
    add_reference_entry(
        doc,
        "S1",
        "Pydantic Models",
        "https://docs.pydantic.dev/latest/concepts/models/",
        "定义实体、关系、实验配置和结构化响应的运行时 validation contract",
    )
    add_reference_entry(
        doc,
        "S2",
        "Instructor structured outputs",
        "https://github.com/567-labs/instructor",
        "以 Pydantic response model 包装模型调用并执行结构化解析与重试",
    )
    add_reference_entry(
        doc,
        "S3",
        "OpenAI Python client",
        "https://github.com/openai/openai-python",
        "连接 OpenAI-compatible HTTP API",
    )
    add_reference_entry(
        doc,
        "S4",
        "Ollama OpenAI compatibility",
        "https://docs.ollama.com/api/openai-compatibility",
        "通过本地 /v1 endpoint 运行 llama3.1:8b 等模型",
    )

    doc.add_heading("D.2 方法与系统设计参考", level=2)
    add_reference_entry(
        doc,
        "S5",
        "Text2KGBench",
        "https://arxiv.org/abs/2308.02357",
        "ontology-driven KG generation、fact extraction 和 conformance 评价背景",
    )
    add_reference_entry(
        doc,
        "S6",
        "Ontology-guided KG Construction from Maintenance Short Texts",
        "https://aclanthology.org/2024.kallm-1.8/",
        "工业短文本的 ontology-guided triplet extraction 与 hallucination 评价",
    )
    add_reference_entry(
        doc,
        "S7",
        "Procedural Knowledge Ontology",
        "https://arxiv.org/abs/2503.20634",
        "Procedure、Action、Tool、Resource 等程序性类别的 schema 参考",
    )
    add_reference_entry(
        doc,
        "S8",
        "Neo4j LLM Graph Builder",
        "https://github.com/neo4j-labs/llm-graph-builder",
        "Property Graph、图导出和 LLM-to-graph 系统设计参考",
    )
    add_reference_entry(
        doc,
        "S9",
        "LlamaIndex Property Graph example",
        "https://developers.llamaindex.ai/python/examples/property_graph/property_graph_basic/",
        "节点、关系与 Property Graph index 的概念参考",
    )
    add_reference_entry(
        doc,
        "S10",
        "LangChain LLMGraphTransformer",
        "https://python.langchain.com/api_reference/experimental/graph_transformers/langchain_experimental.graph_transformers.llm.LLMGraphTransformer.html",
        "allowed nodes、allowed relationships 与 strict filtering 的设计参考",
    )
    add_reference_entry(
        doc,
        "S11",
        "GLiREL",
        "https://github.com/jackboyla/GLiREL",
        "候选关系分类/scorer 接口参考；当前项目尚未完成真实 GLiREL 对照实验",
    )
    add_reference_entry(
        doc,
        "S16",
        "An Approach for Systematic Decomposition of Complex LLM Tasks",
        "https://arxiv.org/abs/2510.07772",
        "Exposé 中 ACONIC-inspired dependency-graph/tree-decomposition 方法的理论来源；当前代码只实现简化 layered extraction",
    )

    doc.add_heading("D.3 数据集与标注来源", level=2)
    add_reference_entry(
        doc,
        "S12",
        "EASG repository",
        "https://github.com/fpv-iplab/EASG",
        "Ego-EASG graph annotations、object grounding 和 metadata adapter",
    )
    add_reference_entry(
        doc,
        "S13",
        "Action Scene Graphs for Long-Form Understanding of Egocentric Videos",
        "https://openaccess.thecvf.com/content/CVPR2024/papers/Rodin_Action_Scene_Graphs_for_Long-Form_Understanding_of_Egocentric_Videos_CVPR_2024_paper.pdf",
        "EASG 人工标注图的研究与数据集说明",
    )
    add_reference_entry(
        doc,
        "S14",
        "IndEgo repository",
        "https://github.com/Vivek9Chavan/IndEgo",
        "工业 egocentric annotation、文本层和 metadata 的来源格式",
    )
    add_reference_entry(
        doc,
        "S15",
        "IndEgo dataset paper",
        "https://papers.nips.cc/paper_files/paper/2025/hash/8e5dc5969a6174fcaaececd890c7f59b-Abstract-Datasets_and_Benchmarks_Track.html",
        "IndEgo 工业场景、任务和 benchmark 背景",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    print(output)
