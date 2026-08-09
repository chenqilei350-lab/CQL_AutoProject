#!/usr/bin/env python3
"""Build a concise Chinese-English code module report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_complete_experiment_report_docx import (
    BLUE,
    INK,
    MUTED,
    ROOT,
    add_body,
    add_bullets,
    add_callout,
    add_code_figure,
    add_figure,
    add_source_note,
    add_table,
    configure_document,
    create_code_figures,
    set_paragraph_border_bottom,
    set_run_font,
)


OUTPUT = ROOT / "docs" / "reports" / "AUT_KG_Simple_Bilingual_Code_Report_2026-08-04.docx"


def add_bilingual(doc: Document, zh: str, en: str) -> None:
    add_body(doc, zh, after=3)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(en)
    set_run_font(run, size=9.5, color=MUTED, italic=True)


def add_bilingual_bullets(doc: Document, pairs: list[tuple[str, str]]) -> None:
    add_bullets(doc, [f"{zh} / {en}" for zh, en in pairs])


def set_report_header(doc: Document) -> None:
    for section in doc.sections:
        paragraph = section.header.paragraphs[0]
        paragraph.text = "AUT KG Extraction Pipeline  |  Bilingual Code Report"
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(paragraph.runs[0], size=8.5, color=MUTED, bold=True)


def build_document() -> Path:
    architecture_chart, code_figures = create_code_figures()
    doc = Document()
    configure_document(doc)
    set_report_header(doc)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run("TU BERLIN · INDUSTRIAL AUTOMATION PROJECT")
    set_run_font(run, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("AUT KG Extraction Pipeline\n简版中英双语代码报告")
    set_run_font(title_run, size=24, color=INK, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle_run = subtitle.add_run("Concise Bilingual Code Module Report")
    set_run_font(subtitle_run, size=13, color=MUTED)

    add_table(
        doc,
        ["项目 / Item", "内容 / Value"],
        [
            ["报告日期 / Date", "2026-08-04"],
            ["语言 / Language", "Python >=3.12"],
            ["环境 / Environment", "uv + Ollama + llama3.1:8b"],
            ["输入 / Input", "UTF-8 industrial .txt or video-derived scene text"],
            ["输出 / Output", "JSON, node/edge CSV, Cypher, Markdown summary"],
        ],
        [2550, 6810],
        font_size=8.7,
        first_column_bold=True,
    )
    rule = doc.add_paragraph()
    set_paragraph_border_bottom(rule, color=BLUE, size=12)
    add_callout(
        doc,
        "当前成品 / Current product",
        "系统已经能够把工业文本一键转换为带节点、关系、证据和关系来源的 Property Graph。The system can convert industrial text into a Property Graph containing nodes, relations, evidence, and relation-origin metadata through a one-click runner.",
    )

    doc.add_page_break()
    doc.add_heading("1. 总体架构 / Overall Architecture", level=1)
    add_figure(
        doc,
        architecture_chart,
        "图 1 / Figure 1. 当前已实现的代码模块与数据流",
        "Implemented modules and data flow of the AUT KG extraction pipeline",
    )
    add_bilingual(
        doc,
        "代码分为两个核心模块。数据清洗模块负责输入适配、确定性清洗、统一结构和证据保留；文本转知识图谱模块负责 schema 约束抽取、逐层抽取、候选关系、图构建、后处理和评价。",
        "The codebase has two core modules. The Data Cleaning Module handles source adaptation, deterministic cleaning, unified representation, and evidence preservation. The Text-to-KG Module handles schema-guided extraction, layered extraction, relation candidates, graph construction, post-processing, and evaluation.",
    )
    add_table(
        doc,
        ["模块 / Module", "主要文件 / Main files", "功能 / Function"],
        [
            ["Cleaning", "backend/cleaning/", "Parse sources; preserve provenance"],
            ["Unified text", "backend/preprocessing/unified_text.py", "Reorganize source-supported facts"],
            ["Schema + LLM", "backend/schemas/; backend/llm/; backend/extraction/", "Typed structured extraction"],
            ["Layered extraction", "backend/pipeline/algorithm_experiment.py", "Entity-first, relation-second"],
            ["Relation candidates", "backend/pipeline/relation_candidate_pipeline.py", "Generate, score, validate candidate edges"],
            ["Graph + postprocess", "backend/graph/; industrial_text_to_kg.py", "Merge, normalize, validate, track origin"],
            ["Evaluation", "backend/evaluation/", "P/R/F1, grounding, hallucination, stability"],
        ],
        [1900, 3300, 4160],
        font_size=7.7,
        first_column_bold=True,
    )

    doc.add_heading("2. 数据清洗与统一文本 / Cleaning and Unified Text", level=1)
    add_code_figure(
        doc,
        code_figures["unified"],
        "代码图 1 / Code Figure 1. UnifiedTextRecord 构建与 evidence 校验",
        "Unified text record construction and evidence validation code",
    )
    add_bilingual(
        doc,
        "数据清洗使用确定性解析器，不调用 LLM 补写事实。UnifiedTextRecord 将 scene、action sequence、tools/objects、parameters、quality result 和 uncertainty 放入固定字段。每条 GroundedEntry 的 evidence 必须能在 raw text 中找到。",
        "Data cleaning uses deterministic parsers and does not ask an LLM to add facts. UnifiedTextRecord places scene, action sequence, tools/objects, parameters, quality results, and uncertainty into fixed fields. Evidence for every GroundedEntry must occur in the raw text.",
    )
    add_bilingual_bullets(
        doc,
        [
            ("输入：来源 JSON/CSV/text layer 或普通工业文本", "Input: source JSON/CSV/text layers or plain industrial text"),
            ("输出：clean records、Unified JSON 和 prompt-ready text", "Output: clean records, Unified JSON, and prompt-ready text"),
            ("边界：只重组原文，不做翻译式润色或事实补充", "Boundary: reorganizes source facts without translation-style rewriting or fact completion"),
            ("实验效果：主要改善 schema、JSON 和 ID 可用性", "Observed effect: mainly improves schema, JSON, and ID usability"),
        ],
    )
    add_source_note(doc, "backend/cleaning/factory_data.py；backend/preprocessing/unified_text.py")

    doc.add_heading("3. Schema 约束与逐层抽取 / Schema-Guided Layered Extraction", level=1)
    add_code_figure(
        doc,
        code_figures["layered"],
        "代码图 2 / Code Figure 2. Entity-first / relation-second 两阶段抽取",
        "Entity first and relation second layered extraction code",
    )
    add_bilingual(
        doc,
        "Pydantic schema 定义 Scene、Action、Tool、SceneObject、Procedure 和 ProcessParameter 等实体，以及 USES_TOOL、ACTS_ON、BEFORE、PART_OF 和 OBSERVED_IN 等关系。Ollama 通过 OpenAI-compatible API 运行本地模型，Instructor 将输出解析为 Pydantic 对象。",
        "The Pydantic schema defines entities such as Scene, Action, Tool, SceneObject, Procedure, and ProcessParameter, together with relations including USES_TOOL, ACTS_ON, BEFORE, PART_OF, and OBSERVED_IN. Ollama runs the local model through an OpenAI-compatible API, while Instructor parses the output into Pydantic objects.",
    )
    add_bilingual(
        doc,
        "逐层模式先生成 entity inventory，再让关系阶段只能引用已经抽取的实体端点。这样可以降低单次任务复杂度并分开诊断节点和关系错误，但第一阶段漏掉的实体会直接限制第二阶段的关系召回。",
        "The layered mode first creates an entity inventory and then restricts the relation stage to those extracted endpoints. This reduces single-call complexity and separates node and relation errors, but entities missed in stage one directly limit relation recall in stage two.",
    )
    add_callout(
        doc,
        "实验解释 / Experimental interpretation",
        "Unified+Layered 达到当前最佳工业 Node F1 0.6019，但最佳 Edge F1 仍来自 Raw+One-shot 0.1421，因此不能声称 layered 已经稳定提高关系质量。Unified+Layered achieved the current best industrial Node F1 of 0.6019, but the best Edge F1, 0.1421, came from Raw+One-shot; layered extraction has not consistently improved relation quality.",
        risk=True,
    )
    add_source_note(doc, "backend/schemas/egocentric_video.py；backend/llm/client.py；backend/extraction/extractor.py；backend/pipeline/algorithm_experiment.py")

    doc.add_heading("4. 关系候选与验证 / Relation Candidates and Validation", level=1)
    add_code_figure(
        doc,
        code_figures["candidate"],
        "代码图 3 / Code Figure 3. RelationCandidateGenerator 候选生成入口",
        "Relation candidate generator entry point",
    )
    add_bilingual(
        doc,
        "关系模块先根据实体类型、文本共现、动作顺序和 annotation metadata 生成有限候选，再由确定性规则或可选 LLM binary judge 判断候选是否受原文支持。模型只能返回 candidate_id、is_supported、confidence、evidence_text 和 reason，不能创造新端点。",
        "The relation module first generates a constrained candidate set from entity types, text co-occurrence, action order, and annotation metadata. Deterministic rules or an optional LLM binary judge then decide whether each candidate is source-supported. The model may return only candidate_id, is_supported, confidence, evidence_text, and reason; it cannot invent endpoints.",
    )
    add_table(
        doc,
        ["Relation", "Candidate rule", "Current status"],
        [
            ["USES_TOOL", "Action -> Tool with local evidence", "Recall bottleneck"],
            ["ACTS_ON", "Action -> Object by verb/object evidence", "Broadly usable"],
            ["BEFORE", "Action sequence or temporal order", "Strong coverage"],
            ["PART_OF", "Action -> Procedure/Keystep", "Partial"],
            ["OBSERVED_IN", "Action -> Scene", "Deterministic"],
            ["WARNING_FOR", "Warning/mistake -> Action/Keystep", "Partial"],
        ],
        [1800, 4500, 3060],
        font_size=8.0,
        first_column_bold=True,
    )
    add_bilingual(
        doc,
        "Binary judge 的 Edge F1 从 0.0476 小幅提高到 0.0564，但平均耗时为 81.84 秒。当前它作为可选 precision filter 保留，主要改进方向仍是提高候选 recall 和 endpoint alignment。",
        "The binary judge increased Edge F1 slightly from 0.0476 to 0.0564 but required 81.84 seconds on average. It remains an optional precision filter; the main priority is still candidate recall and endpoint alignment.",
    )
    add_source_note(doc, "backend/pipeline/relation_candidate_pipeline.py；docs/results/relation_candidate_ablation_2026-07-07.md")

    doc.add_heading("5. 图构建、后处理与评价 / Graph, Post-Processing and Evaluation", level=1)
    add_table(
        doc,
        ["组件 / Component", "实现 / Implementation", "作用 / Purpose"],
        [
            ["Property Graph", "GraphNode, GraphEdge, stable IDs", "Typed in-memory graph"],
            ["Node merge", "Normalized name + token overlap", "Merge repeated entities"],
            ["Relation merge", "(source, relation, target) deduplication", "Remove duplicate edges"],
            ["Direction validation", "Schema domain/range", "Detect or safely reverse wrong edges"],
            ["Grounding", "Evidence checked against source text", "Flag unsupported facts"],
            ["Relation origin", "llm_extracted/fallback/postprocessed", "Separate model and rule contributions"],
            ["Evaluation", "Node/Edge P/R/F1 + Jaccard stability", "Measure quality and repeated agreement"],
        ],
        [2200, 3800, 3360],
        font_size=7.9,
        first_column_bold=True,
    )
    add_bilingual(
        doc,
        "后处理只执行可确定的节点别名、工具分类、关系方向和重复边修正。Raw metrics 与 normalized metrics 分开保存，因此产品稳定输出不会被误写成 LLM 原始能力。",
        "Post-processing applies only deterministic fixes for node aliases, tool classification, relation direction, and duplicate edges. Raw and normalized metrics are stored separately so stable product output is not misreported as raw LLM capability.",
    )
    add_bilingual_bullets(
        doc,
        [
            ("Precision/Recall/F1：正确性和完整性", "Precision/Recall/F1: correctness and completeness"),
            ("Ontology validation：关系 label、domain 和 range", "Ontology validation: relation labels, domains, and ranges"),
            ("Hallucination validation：检查无原文证据事实", "Hallucination validation: detect unsupported facts"),
            ("Stability：重复运行节点和边的 Jaccard overlap", "Stability: repeated-run Jaccard overlap for nodes and edges"),
        ],
    )
    add_source_note(doc, "backend/graph/property_graph.py；backend/pipeline/industrial_text_to_kg.py；backend/evaluation/")

    doc.add_heading("6. Exposé 原方法与当前代码 / Planned vs Implemented Method", level=1)
    add_table(
        doc,
        ["阶段 / Stage", "Exposé 原计划 / Planned", "当前代码 / Implemented", "状态 / Status"],
        [
            ["WP3", "Procedural dependency graph", "Action sequence, temporal metadata, relation candidates", "In progress"],
            ["WP4", "Tree/local-subtask decomposition", "Entity-first/relation-second extraction", "In progress"],
            ["Local extraction", "Independent local ontology-guided subgraphs", "Scene-level one-shot/layered extraction", "Simplified done"],
            ["WP5", "Cross-subgraph entity/conflict merge", "Single-graph node normalization and edge deduplication", "In progress"],
            ["Evaluation", "P/R/F1 and graph consistency", "P/R/F1, grounding, hallucination, stability", "Core done"],
        ],
        [1150, 2900, 3450, 1860],
        font_size=7.6,
        first_column_bold=True,
    )
    add_bilingual(
        doc,
        "调整原因是当前数据主要为短 scene-level 文本，缺少稳定的 dependency labels 和局部子任务边界；本地 7B/8B 模型对完整图生成也容易漏抽或超时。因此项目先实现可运行、可评价、可追踪证据的简化流程，再扩展完整依赖图和跨子图合并。",
        "The method changed because the current data mainly consists of short scene-level text without reliable dependency labels or local-subtask boundaries. Local 7B/8B models also omit facts or time out on full-graph generation. The project therefore implemented a runnable, evaluable, evidence-traceable pipeline before extending it to full dependency graphs and cross-subgraph merging.",
    )
    add_callout(
        doc,
        "论文表述 / Thesis wording",
        "当前方法应称为 simplified layered extraction and evidence-constrained relation validation，而不是完整 ACONIC tree decomposition。The current method should be described as simplified layered extraction and evidence-constrained relation validation, not as a completed ACONIC tree-decomposition pipeline.",
        risk=True,
    )

    doc.add_heading("7. 一键运行与输出 / One-Click Run and Outputs", level=1)
    add_bilingual_bullets(
        doc,
        [
            ("把 UTF-8 工业 .txt 放入 input_texts/", "Place UTF-8 industrial .txt files in input_texts/"),
            ("双击 RUN_TEXT_TO_KG.command", "Double-click RUN_TEXT_TO_KG.command"),
            ("在 output_kg/<input_name>/ 查看结果", "Open results in output_kg/<input_name>/"),
        ],
    )
    add_table(
        doc,
        ["输出文件 / Output", "用途 / Purpose"],
        [
            ["summary.md / RUN_SUMMARY.md", "Per-file summary and batch-level success/error index"],
            ["kg_result.json", "Complete graph, evidence, validation, relation origins"],
            ["graph_nodes.csv / graph_edges.csv", "Node and edge tables for inspection"],
            ["graph.cypher", "Property-graph import/export statements"],
        ],
        [2700, 6660],
        font_size=8.3,
        first_column_bold=True,
    )
    add_source_note(doc, "handoff_packages/industrial_text_to_kg_product_release_2026-07-08_final/")

    doc.add_heading("8. 当前状态与限制 / Current Status and Limitations", level=1)
    add_table(
        doc,
        ["状态 / Status", "内容 / Item"],
        [
            ["Done", "Cleaning, Unified text, Pydantic extraction, Property Graph, outputs, core evaluation"],
            ["In progress", "WP3 dependency model, WP4 decomposition, WP5 cross-subgraph merge, relation recall"],
            ["Not started", "Real GLiREL comparison, direct raw-video processing"],
        ],
        [1900, 7460],
        font_size=8.3,
        first_column_bold=True,
    )
    add_bilingual(
        doc,
        "当前系统已经是可运行的 Product V1，但关系质量仍低于节点质量。下一步优先提高 USES_TOOL 和跨句关系候选召回，并在 32 个工业 scenes 上执行完整端到端重复实验。",
        "The current system is a runnable Product V1, but relation quality remains weaker than node quality. The next priority is improving USES_TOOL and cross-sentence candidate recall, followed by repeated end-to-end evaluation on 32 industrial scenes.",
    )

    doc.add_heading("8.1 方法与代码来源 / Method and Code Sources", level=2)
    add_bilingual_bullets(
        doc,
        [
            ("直接 API：Pydantic、Instructor、OpenAI Python client、Ollama", "Direct APIs: Pydantic, Instructor, OpenAI Python client, and Ollama"),
            ("方法参考：Text2KGBench、KaLLM maintenance KG、PKO、ACONIC decomposition", "Method references: Text2KGBench, KaLLM maintenance KG, PKO, and ACONIC decomposition"),
            ("项目实现：统一文本 evidence 校验、工业 schema、关系候选、后处理、评估和一键入口", "Project implementation: unified-text evidence validation, industrial schema, relation candidates, post-processing, evaluation, and one-click runner"),
        ],
    )

    doc.core_properties.title = "AUT KG Extraction Pipeline - Simple Bilingual Code Report"
    doc.core_properties.subject = "Chinese-English overview of implemented code modules and boundaries"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
